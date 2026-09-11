from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


SOURCE = Path(
    "/Volumes/out/Projects/开题报告/最终选择定版/"
    "开题报告_红壤设施大棚_智能管控平台与根区MPC版.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_红壤设施大棚_智能管控平台与根区MPC版_研究路线修订稿.docx"
)
FONT_PATH = Path(
    "/System/Library/Fonts/PingFang.ttc"
)
if not FONT_PATH.exists():
    FONT_PATH = Path(
        "/System/Library/AssetsV2/com_apple_MobileAsset_Font8/"
        "86ba2c91f017a3749571a82f2c6d890ac7ffb2fb.asset/AssetData/PingFang.ttc"
    )


def set_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            paragraph._p.remove(run._r)
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text: str) -> None:
    set_text(cell.paragraphs[0], text)


def iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def normalize_document_fonts(doc: Document) -> None:
    for style_name in ("Normal", "Body Text", "Title"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    for paragraph in iter_paragraphs(doc):
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            rfonts.set(qn("w:ascii"), "Times New Roman")
            rfonts.set(qn("w:hAnsi"), "Times New Roman")
            rfonts.set(qn("w:eastAsia"), "Songti SC")


def font(size: int, bold: bool = False):
    try:
        return ImageFont.truetype(str(FONT_PATH), size=size, index=1 if bold else 0)
    except OSError:
        return ImageFont.truetype(str(FONT_PATH), size=size)


def wrapped(draw: ImageDraw.ImageDraw, text: str, fnt, max_width: int) -> list[str]:
    lines: list[str] = []
    for block in text.split("\n"):
        line = ""
        for ch in block:
            candidate = line + ch
            if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
                line = candidate
            else:
                if line:
                    lines.append(line)
                line = ch
        lines.append(line)
    return lines


def box(draw, xy, text, fill="#EEF4F8", outline="#6B7C8C", title=False):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=4)
    fnt = font(54 if title else 46, bold=title)
    lines = wrapped(draw, text, fnt, x2 - x1 - 50)
    line_h = int(fnt.size * 1.35)
    total = line_h * len(lines)
    y = y1 + (y2 - y1 - total) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=fnt, fill="#1F2933")
        y += line_h


def arrow(draw, start, end, fill="#526574", width=8):
    draw.line([start, end], fill=fill, width=width)
    x, y = end
    sx, sy = start
    if abs(x - sx) > abs(y - sy):
        direction = 1 if x > sx else -1
        pts = [(x, y), (x - 24 * direction, y - 16), (x - 24 * direction, y + 16)]
    else:
        direction = 1 if y > sy else -1
        pts = [(x, y), (x - 16, y - 24 * direction), (x + 16, y - 24 * direction)]
    draw.polygon(pts, fill=fill)


def make_architecture(path: Path) -> None:
    im = Image.new("RGB", (2362, 1921), "white")
    d = ImageDraw.Draw(im)
    box(d, (330, 90, 2030, 310), "管理与展示层\n棚室配置  任务审核  风险状态  结果追溯", title=True)
    box(d, (820, 530, 1540, 930), "任务与控制编排\n版本冻结  准入检查\n联锁许可  命令反馈", title=True)
    box(d, (120, 530, 700, 930), "Python算法服务\n数据准入  双层灰箱\n强规则  区间MPC\n冻结评价  可信监督")
    box(d, (1660, 530, 2240, 930), "证据与状态存储\nSQLite审计账\n数据 模型 配置哈希\n任务 命令  反馈  事件")
    box(d, (330, 1130, 2030, 1370), "通信与执行接口  MQTT 3.1.1 / REST\n消息幂等  乱序缓冲  超时推进  断连重连", title=True)
    box(d, (330, 1580, 2030, 1810), "当前 仿真对象与本机软件在环\n条件具备后  双深度水分 气象 流量 泵阀与EC pH监测", title=True)
    arrow(d, (1180, 310), (1180, 530))
    arrow(d, (700, 730), (820, 730))
    arrow(d, (1540, 730), (1660, 730))
    arrow(d, (1180, 930), (1180, 1130))
    arrow(d, (1180, 1370), (1180, 1580))
    arrow(d, (1090, 1580), (1090, 1370))
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path, dpi=(300, 300))


def make_validation(path: Path) -> None:
    im = Image.new("RGB", (817, 640), "white")
    d = ImageDraw.Draw(im)
    small = font(19)
    bold = font(21, bold=True)

    def sbox(x1, y1, x2, y2, text, fill="#EEF4F8"):
        d.rounded_rectangle((x1, y1, x2, y2), radius=8, fill=fill, outline="#6B7C8C", width=2)
        lines = wrapped(d, text, bold if y2 - y1 > 55 else small, x2 - x1 - 18)
        lh = 25
        y = y1 + (y2 - y1 - lh * len(lines)) / 2
        for line in lines:
            bb = d.textbbox((0, 0), line, font=bold if y2 - y1 > 55 else small)
            d.text((x1 + (x2 - x1 - (bb[2] - bb[0])) / 2, y), line, font=bold if y2 - y1 > 55 else small, fill="#1F2933")
            y += lh

    sbox(120, 25, 697, 90, "数据质量与辨识信息两级准入\n不充分则补采或缩减模型", "#E6F0F7")
    sbox(120, 120, 697, 185, "双层灰箱辨识与独立循环验证")
    sbox(70, 220, 367, 285, "开发集\n选择强规则与MPC候选")
    sbox(450, 220, 747, 285, "锁定前冻结\n配置 场景 指标与分类门")
    sbox(120, 320, 697, 390, "独立锁定场景配对评价\n灌水 亏缺 过湿 安全 动作次数", "#E6F0F7")
    sbox(70, 430, 367, 505, "异常监督机制检验\n保留失败与停止结论")
    sbox(450, 430, 747, 505, "平台软件在环\n任务 消息 反馈与审计")
    sbox(120, 550, 697, 620, "条件触发升级\n真实回放 影子控制 低风险实体试验", "#FFF3DF")
    arrow(d, (408, 90), (408, 120), width=4)
    arrow(d, (408, 185), (218, 220), width=4)
    arrow(d, (408, 185), (598, 220), width=4)
    arrow(d, (218, 285), (330, 320), width=4)
    arrow(d, (598, 285), (486, 320), width=4)
    arrow(d, (408, 390), (218, 430), width=4)
    arrow(d, (408, 390), (598, 430), width=4)
    arrow(d, (218, 505), (340, 550), width=4)
    arrow(d, (598, 505), (476, 550), width=4)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path, dpi=(300, 300))


def rebuild_with_media(docx_path: Path, replacements: dict[str, Path]) -> None:
    tmp = docx_path.with_suffix(".media.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in replacements:
                zout.writestr(item, replacements[item.filename].read_bytes())
            else:
                zout.writestr(item, zin.read(item.filename))
    tmp.replace(docx_path)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)

    set_text(
        doc.paragraphs[6],
        "论文题目：面向红壤区设施大棚的智慧管控平台设计与根区水分MPC适用性研究",
    )

    t = doc.tables[0].rows[0].cells[0].paragraphs
    set_text(t[2], "设施大棚水肥管理需要连续监测土壤、气象与设备状态，并按分区组织灌溉任务。智慧农业政策及水肥一体化设备、微灌工程和设施园艺物联网标准，对信息感知、设备互联和精准作业提出了要求[1-4]。具体到棚内根区，水分变化同时受灌水、蒸散、层间交换、深层渗漏和执行偏差影响。固定时长或单阈值灌溉便于实施，却难以在未来扰动、泵阀能力和含水率安全区间之间协调，异常情况下也缺少可追溯的降级依据。")
    set_text(t[3], "本课题拟设计面向红壤区设施大棚的智慧管控平台，以根区水分为当前可验证的控制对象。平台统一接入双深度水分、棚内气象、灌水命令、实际流量及设备状态，管理数据准入、模型版本、控制任务、安全联锁和执行反馈。肥液EC、pH、压力、液位等作为监测与扩展接口保留；在缺少养分对象数据和实体设备的阶段，不把水分控制结果外推为水肥协同优化效果。")
    set_text(t[5], "根区水分模型预测控制（Model Predictive Control，MPC）作为平台灌溉决策的核心算法。研究采用10 cm与25 cm双深度体积含水率描述浅层快速响应和深层蓄水变化，以实际到达水量而非仅以命令水量驱动模型。MPC根据当前状态、短时气象扰动和动作约束滚动计算候选灌水量；数据或模型未通过准入时不发布MPC版本，观测、求解或执行异常时进入规则回退或安全暂停。")
    set_text(t[6], "研究对象定位于红壤土壤栽培大棚，但当前尚无红壤连续实测数据。现阶段使用独立合成对象验证数据准入、双层辨识、冻结控制评价和平台消息链，结论仅表述为方法与软件在环证据。后续须通过同一设备和工况下的传感器标定、流量标定及湿润—干燥循环数据重新辨识红壤参数；模型通过独立验证后，方可进入真实数据回放或低风险现场试验。")
    set_text(t[8], "本课题把平台建设与算法研究组织为一条证据链：数据先接受质量和辨识信息准入，模型在独立循环验证后才能驱动控制器，规则与MPC在开发—锁定隔离的场景中比较，异常监督接受独立失效检验，最终由任务—版本—命令—反馈审计还原全过程。该路线既保留智慧管控平台主体，也直接回应灌溉MPC受现场数据不足、模型失配和执行异常制约的问题。")

    t = doc.tables[1].rows[0].cells[0].paragraphs
    set_text(t[2], "设施水肥研究已由单机配肥扩展到环境感知、网络通信、远程管理和决策协同。国内相关工作覆盖水肥装备、温室环境监测、人机交互及物联网联动[6-12]，国外研究也形成营养液混配、自动灌溉施肥和滴灌动态仿真工具[13-15]。现有成果为设备与信息系统集成提供了基础，但对算法进入控制前的数据准入、模型版本冻结、命令与实际供水核对以及异常消息后的状态恢复，仍需要统一的任务化证据链。")
    set_text(t[4], "MPC使用过程模型预测未来响应，并在每个控制周期滚动求解约束优化问题[23]。灌溉研究已探索数据驱动MPC[24]、天气预报不确定性下的根区亏缺与灌水量优化[25]、农业水文对象的区间MPC[26]以及考虑水能成本的周期MPC[27]。这些工作表明，MPC用于灌溉并非空白；其研究价值主要来自对根区动态、未来扰动和动作约束的统一处理。相应地，本课题不以“首次将MPC用于灌溉”为创新主张，而聚焦有限数据下模型何时可用、控制价值如何冻结评价以及失败如何被平台记录。")
    set_text(t[5], "灌溉MPC能否落地首先受控制导向模型约束。采样点数量或湿干循环数量不能证明参数已被充分激励，训练误差较小也不能替代独立循环预测、物理边界和闭环价值检验。红壤参数、根区深度、蒸散估计、传感标定和实际灌水量的变化都会形成模型—对象失配。本课题因此设置结构质量与幅值敏感辨识信息两级准入，并把模型辨识、控制器开发和锁定评价使用的数据与场景相互隔离。")
    set_text(t[7], "网络化控制和边缘计算研究为现场分层计算及时延、丢包分析提供了基础[28-30]；故障诊断与容错控制研究则表明，检测、隔离、控制重构和恢复需要协同设计[31-35]。对大棚灌溉而言，异常检测并不自动产生安全回退：传感高读、传感低读、流量衰减和阀门卡滞需要的动作方向可能相反，观测隔离还可能阻断恢复。平台必须把风险状态、允许动作、恢复证据和硬联锁分别记录，并以亏缺、过湿和安全违反检验回退效果。")
    set_text(t[9], "现有研究已经证明灌溉MPC具有理论与应用基础，当前缺口不宜概括为“算法发展不足”，而应具体表述为三点：有限现场数据是否含有足够辨识信息；不同模型与接口能否形成一致的闭环证据；异常检测能否转化为可达且安全的回退与恢复动作。本课题以智慧管控平台为载体，把数据准入、双层灰箱模型、冻结配对评价、监督失效分析和命令反馈审计纳入同一研究流程。")

    t = doc.tables[2].rows[0].cells[0].paragraphs
    set_text(t[2], "设计并实现面向红壤区设施大棚的智慧管控平台原型，形成多源数据接入、数据与模型准入、灌溉任务、控制模式、安全联锁、命令反馈和证据追溯能力。建立双深度根区灰箱模型与区间MPC，在独立合成对象上以强规则为基线开展冻结配对评价，并通过异常场景研究回退与恢复的有效条件。取得同源现场数据后，再进行红壤参数辨识、真实回放和低风险实体升级。")
    set_text(t[3], "平台总体架构如图1所示。当前正式算法与软件在环证据由Python工程统一生成，SQLite保存本机审计状态，MQTT/REST承担消息与服务接口；管理界面和生产数据库可在现场需求明确后扩展。界面或外部服务不能绕过数据准入、版本冻结和安全联锁直接生成泵阀指令。")
    set_text(t[6], "（1）需求分析与平台架构：明确棚室—分区—作物—土壤—设备—算法—任务的数据关系，设计数据接入、算法服务、任务编排、联锁、消息和审计接口。当前以Python、SQLite和MQTT形成可运行内核，交付需求规格、架构图、数据字典、接口协议和软件在环原型。")
    set_text(t[7], "（2）数据准入与双层根区模型：建立双深度水分、气象、命令灌水、实际灌水和流量的统一时间语义；完成结构质量、传感与流量标定、幅值敏感信息量、预测不确定性及用途隔离检查。准入后拟合双层灰箱模型，并在独立循环上验证单步与整循环预测、偏差、物理边界和参数约束。")
    set_text(t[8], "（3）根区水分MPC与可信监督：建立由区间亏缺、过湿、灌水投入和动作变化构成的控制目标，设置状态、单次水量和执行能力约束。使用开发场景选择规则与MPC候选，冻结后进入锁定场景；异常监督分别评价风险检测、规则回退、安全暂停和恢复可达性，未通过晋级门的版本保留为失效机理证据。")
    set_text(t[9], "（4）平台功能与可追溯运行：实现任务、数据、模型、控制器、监督器、命令、反馈、联锁和模式事件的关联；检验重复、乱序、超时、进程异常、MQTT断连重连和异常时间戳。EC、pH和施肥命令仅保留扩展接口，取得养分数据前不纳入已验证MPC对象。")
    set_text(t[10], "（5）系统验证与条件升级：先完成烟雾试验，再开展开发选择和独立锁定确认；主要比较强规则与MPC的灌水量、双层加权亏缺、过湿、安全违反和动作次数。软件在环验证平台语义；现场条件满足后按“数据回放—影子控制—低风险短周期闭环”逐级升级。")
    set_text(t[12], "（1）如何判断有限现场数据不仅格式完整，而且对双层根区模型提供了足够的辨识信息。")
    set_text(t[13], "（2）如何在模型、状态、时间步和灌水反馈接口一致的条件下，冻结评价MPC相对强规则的联合价值与动作代价。")
    set_text(t[14], "（3）如何把异常检测转换为可达的回退与恢复动作，并识别安全预算在不同故障类型下的冲突。")
    set_text(t[15], "（4）如何用统一任务和版本标识连接数据准入、控制建议、联锁许可、设备反馈及论文证据，使正常与失败结果均可复核。")

    rows = doc.tables[3].rows
    t = rows[0].cells[0].paragraphs
    set_text(t[2], "研究按照“平台边界—数据准入—模型辨识—控制开发—锁定评价—异常监督—软件审计—条件升级”推进。现阶段先用独立合成对象和本机消息链验证方法，禁止使用锁定场景反向调参；取得现场数据后沿同一数据结构重新执行准入与验证，而不直接替换合成参数并宣称红壤有效。")
    set_text(t[3], "2. 双层根区灰箱模型与数据准入")
    set_text(t[4], "以10 cm与25 cm体积含水率组成状态向量x(k)=[θ10(k), θ25(k)]ᵀ。离散灰箱模型分别描述实际灌水对浅层的增益、太阳辐射与饱和水汽压差引起的耗水、浅深层交换及深层排水。第k行水分表示区间末状态，第k行实际灌水和气象量表示前一采样区间输入，从而避免状态与输入错位。参数来源、数据用途、配置和代码哈希随模型版本保存。")
    set_text(t[5], "数据先检查字段、时间、缺失、物理范围、湿润—干燥动态和水量账，再依据标定误差、幅值敏感信息量和预测不确定性输出“数据质量阻塞”“继续补采激励”或“暂准入辨识”。模型仅使用辨识循环拟合，验证循环用于冻结检查。若安全范围内无法获得足够激励，则固定弱可辨识参数或缩减模型结构，不以增加采样行数替代信息充分性。")
    set_text(t[7], "MPC在每个15 min控制周期读取双层状态、未来气象序列和执行能力，在预测时域内权衡浅深层目标区间、灌水投入、过湿风险和动作变化。优化输出候选命令，平台联锁根据观测有效性、执行器状态、流量反馈和消息状态形成最终命令。规则与MPC共享相同的状态、时间和命令—反馈接口。")
    set_text(t[8], "监督器把观测风险、模型风险和执行风险分开计算，并维护MPC、规则回退与安全暂停模式。版本晋级要求异常窗口内的亏缺与过湿风险改善且不增加安全违反；若检测触发后无法恢复，或安全预算对某类执行故障放大缺水尾部，则停止继续调参并记录机理。回退策略在独立证据通过前不表述为已经保证安全。")
    set_text(t[9], "灌溉决策与执行流程如图2所示。模型或求解异常且基本观测有效时可进入规则回退；关键观测或执行条件失效时进入安全暂停。所有候选动作均经联锁，反馈以任务标识和逻辑序号核对；迟到、重复、乱序或冲突反馈保留审计记录，不能重复改变已经确定的控制状态。")
    set_text(t[11], "4. 平台实现方法")
    set_text(t[12], "当前以纯Python实现数据清洗与准入、双层灰箱辨识、强规则与MPC、bootstrap统计、监督状态机、SQLite事务审计和MQTT代理/客户端试验。正式结果由统一虚拟环境、配置和脚本生成，MATLAB/Simulink仅作为导师要求时的框图展示或独立数值复核工具。后续若建设完整管理界面，可按需求增加Java Spring Boot、MySQL和Vue，但其完成度不作为算法结论成立的前置条件。")
    set_text(t[13], "平台区分离线模型生命周期与在线任务生命周期。离线流程冻结原始数据、用途、模型和控制器版本；在线流程按观测、风险、模式、候选动作、联锁命令和反馈顺序推进。MQTT 3.1.1与QoS 1负责消息传输[36]，SQLite状态机负责应用层幂等、顺序、超时和恢复。模型更新必须创建新版本和新任务，不能覆盖已产生的控制证据。")
    set_text(t[15], "设置强规则与双层MPC两类主对照，可信监督作为独立机制试验，不预设其必然优于基础MPC。场景采用相同天气、初始浅深水分、参数扰动和随机种子配对；开发场景用于候选选择，锁定场景只在配置冻结后运行。异常场景覆盖传感偏置、漂移、冻结、反馈冲突、流量衰减、阀门卡滞、通信中断和求解失败。")
    set_text(t[16], "主要终点为命令与实际灌水量、浅深层亏缺及其加权值、过湿积分和安全违反；同时报告动作变化次数、求解失败、模式切换、恢复时长和消息一致性。统计报告场景级配对中位差、bootstrap区间、最小差值和联合不劣比例。没有独立流量计量、完整生育期和重复小区时，不评价真实节水率、肥料利用率、产量或品质。")
    set_text(t[17], "分阶段验证流程如图3所示。数据与模型未通过准入时停止在补采或缩减结构；控制器只在开发集选择后冻结；锁定评价不得用于回改权重；监督版本未超过预设机制门时不进入确认集。软件在环只支持算法流程和平台语义，现场效果须由真实回放、影子控制或具备联锁与人工接管的低风险实体试验补充。")

    t = rows[1].cells[0].paragraphs
    set_text(t[1], "（1）提出面向有限根区数据的两级准入方法，将结构质量、标定误差与幅值敏感辨识信息结合，使平台能够输出补采、缩减模型或准入决定。")
    set_text(t[2], "（2）建立通过准入的双层灰箱模型到区间MPC的统一接口，并采用强规则基线、开发—锁定隔离和多终点配对统计评价灌水—亏缺联合价值及动作代价。")
    set_text(t[3], "（3）把可信监督作为可证伪的机制研究，分别检验风险触发、回退动作、恢复可达性和安全预算，保留恢复死锁与故障类型冲突等否定性证据。")
    set_text(t[4], "（4）将数据、模型、控制模式、联锁、命令与反馈绑定到统一任务和版本审计链，使平台既承担灌溉管理，也成为算法比较与结论复核的研究基础设施。")

    t = rows[2].cells[0].paragraphs
    set_text(t[2], "已形成可重复运行的Python研究工程，覆盖数据质量与模型准入、双层灰箱辨识、强规则与区间MPC、异常监督、SQLite审计及MQTT软件在环；当前99项自动测试通过。前期90个独立锁定合成场景中，双层MPC相对强规则的命令灌水配对中位差为−6.5 mm，加权亏缺配对中位差为−0.000314，联合不劣比例为75.56%，两方法均无安全违反；但MPC动作变化配对中位增加130次，任一主要终点不利比例24.44%接近25%预设上限。上述结果仅作为方法可运行和后续研究风险识别依据，不代表红壤现场节水效果。")
    set_text(t[3], "研究者具备Java、Python、MATLAB、数据库和前端开发基础。当前只有仿真与本机软件在环条件，尚缺合作大棚连续数据、红壤实测参数、实体传感器接入和完整生育期作物试验。因此，毕业设计最低成果定位为可审计平台原型、有限数据准入方法、双层MPC冻结仿真评价及监督失效边界；现场结果属于条件触发的增强层。")
    set_text(t[5], "（1）缺少红壤与作物参数：取得数据前仅使用独立合成对象验证方法，不把其命名为红壤数字孪生；后续优先选定一个土壤栽培棚和一种代表性作物，完成土壤取样、双深度水分与流量标定。")
    set_text(t[6], "（2）连续数据量不足：不预设固定样本量能够保证研究成功。先要求至少3个辨识循环和2个冻结验证循环作为用途隔离入口，再按幅值敏感信息量、预测不确定性和独立验证逐循环决定补采、缩减模型或停止。")
    set_text(t[7], "（3）传感与执行误差混入模型：统一记录命令水量、实际流量、累计水量和区间结束时间；测量误差与动态残差尺度不相容时拒绝生成伪过程噪声参数，待同源标定数据补齐后重估。")
    set_text(t[8], "（4）MPC动作过于频繁：将动作变化次数、最小开闭时间和单周期灌水量纳入正式指标；在锁定结果不回调的前提下开展执行器友好候选开发，必要时保持强规则为部署方案。")
    set_text(t[9], "（5）可信回退难以同时处理多类故障：避免无边界开发新版本。优先补充独立观测或恢复探测通道，并按异常类型区分传感故障与执行故障的预算和恢复条件；新机制未通过开发晋级门时不使用锁定集。")
    set_text(t[10], "（6）缺少实体条件：先完成真实格式数据入口、消息压力试验和影子控制接口；只有流量反馈、硬联锁、人工接管和短周期安全协议具备后，才进入低风险闭环。若毕业前仍无现场条件，论文严格限定为仿真与软件在环适用性研究。")

    schedule = doc.tables[4]
    schedule_updates = {
        2: ("冻结课题边界、双层端到端协议和软件在环基线，完成开题报告修订。", "开题报告、协议、任务台账、初始原型"),
        3: ("完善数据模板、两级准入、模型版本和任务—命令—反馈审计，整理实时核验文献。", "数据字典、准入报告、证据审计、最小工作台"),
        4: ("争取同源红壤标定与湿干循环数据；无现场条件时完成受控合成对象和公开气象回放。", "标定记录或数据缺口报告、模型准入判定"),
        5: ("准入后辨识双层灰箱模型并独立验证；完成执行器友好MPC候选开发。", "冻结模型、控制器、独立验证报告"),
        6: ("开展强规则与MPC锁定配对评价，复核动作代价；检验异常监督的新恢复通道。", "确认性结果、失效机理与停止结论"),
        7: ("完善Python平台内核、MQTT/REST接口和审计展示；条件允许时接入真实数据影子运行。", "可运行平台、接口与消息压力报告"),
        8: ("完成综合复核；具备联锁与人工接管时增加低风险实体试验，否则冻结软件在环边界。", "可复现实验包、条件验证或边界报告"),
    }
    for row_i, (content, outcome) in schedule_updates.items():
        set_cell_text(schedule.rows[row_i].cells[1], content)
        set_cell_text(schedule.rows[row_i].cells[2], outcome)

    refs = schedule.rows[15].cells[0].paragraphs
    set_text(refs[0], "八、主要参考文献（共36项）")
    set_text(refs[23], "[23] QIN S J, BADGWELL T A. A survey of industrial model predictive control technology[J]. Control Engineering Practice, 2003, 11(7): 733-764. DOI:10.1016/S0967-0661(02)00186-7.")
    set_text(refs[24], "[24] BWAMBALE E, ABAGALE F K, ANORNU G K. Data-driven model predictive control for precision irrigation management[J]. Smart Agricultural Technology, 2023, 3: 100074. DOI:10.1016/j.atech.2022.100074.")
    set_text(refs[25], "[25] DELGODA D, MALANO H, SALEEM S K, et al. Irrigation control based on model predictive control (MPC): Formulation of theory and validation using weather forecast data and AQUACROP model[J]. Environmental Modelling & Software, 2016, 78: 40-53. DOI:10.1016/j.envsoft.2015.12.012.")
    set_text(refs[26], "[26] MAO Y, LIU S, NAHAR J, et al. Soil moisture regulation of agro-hydrological systems using zone model predictive control[J]. Computers and Electronics in Agriculture, 2018, 154: 239-247. DOI:10.1016/j.compag.2018.09.011.")
    set_text(refs[27], "[27] CÁCERES G, MILLÁN P, PEREIRA M, et al. Smart Farm Irrigation: Model Predictive Control for Economic Optimal Irrigation in Agriculture[J]. Agronomy, 2021, 11(9): 1810. DOI:10.3390/agronomy11091810.")

    outline = schedule.rows[16].cells[0].paragraphs
    set_text(outline[1], "第1章 绪论：工程背景、研究现状、有限数据问题、研究内容和技术路线。")
    set_text(outline[2], "第2章 系统需求与总体架构：应用边界、双层根区对象、离线模型生命周期、在线控制反馈和证据审计。")
    set_text(outline[3], "第3章 数据准入与双层根区灰箱模型：标定、时间对齐、辨识信息、模型方程及独立循环验证。")
    set_text(outline[4], "第4章 强规则与双层根区MPC冻结评价：候选开发、锁定协议、多终点配对结果和动作代价。")
    set_text(outline[5], "第5章 可信监督与失效机理：风险检测、规则回退、安全暂停、恢复死锁、预算冲突和停止规则。")
    set_text(outline[6], "第6章 智慧管控平台软件在环验证：任务账目、可靠消息、MQTT、SQLite状态恢复及证据审计。")
    set_text(outline[7], "第7章 综合讨论与条件触发路线：仿真证据、红壤模型、实体效果和工程部署边界。")
    set_text(outline[8], "第8章 结论与展望：数据准入、控制评价、监督机制、平台贡献及后续现场验证。")

    normalize_document_fonts(doc)

    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        architecture = td_path / "image2.png"
        validation = td_path / "image4.png"
        make_architecture(architecture)
        make_validation(validation)
        doc.save(OUTPUT)
        rebuild_with_media(
            OUTPUT,
            {
                "word/media/image2.png": architecture,
                "word/media/image4.png": validation,
            },
        )
    print(OUTPUT)


if __name__ == "__main__":
    main()
