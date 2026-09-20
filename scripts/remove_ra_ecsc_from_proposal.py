from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from redesign_proposal_technical_route import make_technical_route
from revise_proposal_from_current_evidence import (
    normalize_document_fonts,
    rebuild_with_media,
    set_text,
)


SOURCE = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与协同控制的智能水肥一体化管控系统设计与实现_以红壤丘陵区为例_最终题名稿.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与协同控制的智能水肥一体化管控系统设计与实现_以红壤丘陵区为例_算法通用表述稿.docx"
)


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main() -> None:
    document = Document(SOURCE)
    replacements = {
        "风险感知自适应边云监督控制（RA-ECSC）": "风险感知协同控制策略",
        "RA-ECSC控制器与冻结配置": "协同控制策略与冻结配置",
        "RA-ECSC": "风险感知协同控制策略",
    }
    for paragraph in all_paragraphs(document):
        original = paragraph.text
        revised = original
        for old, new in replacements.items():
            revised = revised.replace(old, new)
        if revised != original:
            set_text(paragraph, revised)

    normalize_document_fonts(document)
    document.save(OUTPUT)

    with tempfile.TemporaryDirectory(prefix="proposal_t059_") as folder:
        route = Path(folder) / "technical_route.png"
        make_technical_route(
            route,
            problem_text="红壤丘陵区应用场景下通信与传感异常并发的水肥一体机EC连续控制",
        )
        rebuild_with_media(OUTPUT, {"word/media/image4.png": route})


if __name__ == "__main__":
    main()
