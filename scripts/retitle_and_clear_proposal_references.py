from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.shared import Pt
from docx.table import _Cell, Table

from revise_proposal_from_current_evidence import normalize_document_fonts, set_text


SOURCE = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于风险感知边云协同控制的智能水肥一体化系统设计与实现.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现_无参考文献版.docx"
)

OLD_TITLE = "基于风险感知边云协同控制的智能水肥一体化系统设计与实现"
NEW_TITLE = "基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现"
OLD_SYSTEM_NAME = "基于风险感知边云协同控制的智能水肥一体化系统"
NEW_SYSTEM_NAME = "基于物联网与边云协同控制的智能水肥一体化管控系统"
CITATION_PATTERN = re.compile(r"\[(?:\d+(?:\s*[-,，]\s*\d+)*)\]")


def iter_paragraphs(parent):
    if isinstance(parent, _Document):
        for paragraph in parent.paragraphs:
            yield paragraph
        for table in parent.tables:
            yield from iter_paragraphs(table)
    elif isinstance(parent, Table):
        seen = set()
        for row in parent.rows:
            for cell in row.cells:
                key = id(cell._tc)
                if key in seen:
                    continue
                seen.add(key)
                yield from iter_paragraphs(cell)
    elif isinstance(parent, _Cell):
        for paragraph in parent.paragraphs:
            yield paragraph
        for table in parent.tables:
            yield from iter_paragraphs(table)


def remove_extra_paragraphs(cell: _Cell) -> None:
    paragraphs = list(cell.paragraphs)
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def main() -> None:
    document = Document(SOURCE)

    for paragraph in iter_paragraphs(document):
        original = paragraph.text
        revised = original.replace(OLD_TITLE, NEW_TITLE)
        revised = revised.replace(OLD_SYSTEM_NAME, NEW_SYSTEM_NAME)
        revised = CITATION_PATTERN.sub("", revised)
        if revised != original:
            set_text(paragraph, revised)

    set_text(document.paragraphs[6], f"论文题目：{NEW_TITLE}")
    for run in document.paragraphs[6].runs:
        run.font.size = Pt(17)

    reference_cell = document.tables[4].rows[15].cells[0]
    set_text(reference_cell.paragraphs[0], "八、主要参考文献")
    remove_extra_paragraphs(reference_cell)

    normalize_document_fonts(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
