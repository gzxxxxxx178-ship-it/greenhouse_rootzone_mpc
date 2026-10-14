from pathlib import Path
from copy import deepcopy
import csv
import json
import re
import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from revise_proposal_from_current_evidence import rebuild_with_media

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('/Volumes/out/Projects/开题报告/开题报告_引用规范修订版.docx')
OUTPUT=Path('/Volumes/out/Projects/开题报告/开题报告_便携式样机与协同控制修订版.docx')
ASSETS=ROOT/'output/researchwrite/t062_portable_prototype'

# Proposals describe planned methods and evidence; no prototype results are assumed.
REPLACEMENTS={
 (0,0,2): '精准农业将感知、数据处理和控制执行连接起来，为水肥装备的自动化管理提供系统基础[1-2]。物联网与农业自动化技术进一步将设备状态采集、远程管理和数据辅助决策纳入水肥作业流程[3-5]。面向红壤丘陵区分散生产单元，本课题选择便携式水肥一体机作为系统载体，研究设备在不同接入点轮换作业时的EC调节与任务管理。水肥一体机以电导率（electrical conductivity，EC）作为混合肥液浓度的主要反馈量；EC模型控制、系统辨识与模糊PID、自适应PI为混肥动态调节提供了方法依据[6-8]。输配过程仿真和营养液配制研究也说明，控制设计应结合管路输配、混合动态及EC、pH监测需求[9-10]。便携设备转场后需要重新连接供水与灌溉管网，因此本研究将供水压力、流量、管路条件和初始EC的变化纳入工况表征，以样机测量建立控制对象与接入条件之间的联系。',
 (0,0,3): '设备—边缘—云连续体为现场响应与远端计算的分工提供了架构基础[11-13]，边缘计算及多层工业控制研究讨论了控制任务靠近现场部署的方法[14-15]。预测补偿和网络容错控制为通信不确定性处理提供了依据[16-17]。便携式水肥一体机在转场接入和连续作业中，需要协调对象条件、消息时效性与本地观测可信度：新任务启动时核对接入参数并建立可信初始状态，同一任务运行中依据通信和观测风险分配控制权。由此，将转场后的任务初始化与作业中的模式切换连接起来，是本课题系统设计与控制方法的共同切入点。',
 (0,0,5): '本课题以红壤丘陵区设施农业为应用场景，拟设计基于物联网与协同控制的智能水肥一体化管控系统，以便携式水肥一体机完成不同接入点的固定作业。系统由EC与设备状态采集、注肥执行接口、边缘控制节点、云端服务、数据库和管理界面构成，采用MQTT 3.1.1组织发布订阅通信[18]。控制层设置M0云端预测、M1边缘控制和M2模型锚定降级三种模式，联合通信风险与观测风险确定控制位置和反馈来源。转场后以新的任务标识绑定接入点、配方、模型与控制参数；连续作业中的模式切换则通过状态同步和无扰衔接保持控制动作平滑。',
 (0,0,6): '研究按“装备接入—控制方法—系统实现—样机验证”展开。首先测量样机供水、注肥和混肥动态，建立接入工况与模型配置；随后设计双风险三模式监督及转场启动管理；再实现MQTT消息链、任务审计和状态展示。验证采用配对仿真、软件在环和样机对照试验，围绕接入条件变化、通信异常、观测异常及重新启动组织证据。',
 (0,0,8): '便携式装备将水肥作业从单一固定接入点扩展到多个生产单元，系统管理也随之涉及接入配置、任务交接与控制恢复。本课题将这些环节纳入统一控制链：以样机动态测量支撑模型配置，以通信和观测双风险支撑控制权分配，以任务标识和事件记录支撑跨地块追溯。研究价值在于把丘陵分散地块的设备使用需求转化为可验证的接入管理与EC控制机制，为便携水肥装备的物联网集成提供方法和实现依据。',
 (1,0,9): '农业数字孪生部署研究将模型服务与云—雾—边基础设施联系起来，为现场系统与多层计算的协作提供参考[36]。移动式果园灌溉施肥机研究已将混肥、检测、自动控制及牵引行走装置集成到样机中，并开展了灌溉与施肥性能试验[37]。本课题在便携装备应用基础上，进一步研究转场接入条件与控制状态的衔接：用样机数据表征不同接入工况，在新任务启动时重新建立可信状态，在同一任务作业中实施双风险三模式监督，并以仿真、软件在环和样机试验评价控制恢复及运行连续性。',
 (2,0,2): '面向红壤丘陵区分散生产单元，设计并实现适配便携式水肥一体机的物联网协同管控系统，形成接入工况配置、混肥对象建模、双风险三模式监督、转场任务初始化、无扰切换与作业追溯的完整流程。研究将以样机测量支撑模型和工况设置，通过冻结的配对仿真、MQTT软件在环及样机对照试验，评价EC调节、接入后恢复和控制执行的一致性。',
 (2,0,3): '系统总体架构由样机感知与执行、边缘控制、MQTT通信、云端服务、数据库和管理界面构成。边缘节点采集EC、供水流量、压力、液位和执行状态，并完成时间标记、接入检查和本地控制；云端管理接入点配置、模型版本、配方与任务。设备转场后建立新任务，连续运行中由监督器选择M0、M1或M2。每条观测、命令、反馈和模式事件均关联设备、接入点、任务及版本。',
 (2,0,5): '（1）便携样机接入与系统总体设计。围绕转场、连接、启动、作业和结束五个环节，确定样机供水、注肥、传感和通信接口；记录设备质量、外形尺寸、搬运方式、连接形式和安装调试时间，形成接入规格与作业流程。平台设置设备和接入点配置，核对配方、管路工况、传感状态及模型版本后建立新任务。',
 (2,0,6): '（2）样机动态测量与混肥对象建模。在样机允许的压力、流量和注肥范围内，开展阶跃与启动响应试验，记录入口水EC、混肥出口EC、温度、压力、流量及注肥动作。以独立试验记录分别承担辨识和验证，建立注肥输入至EC偏差的一阶惯性纯滞后模型，分析接入工况对增益、时间常数和滞后的影响。各方法共享离散PI、抗积分饱和与执行器变化率约束。',
 (2,0,7): '（3）双风险监督与转场启动管理。以命令年龄、时延EWMA和丢包率表征通信风险，以双EC通道差和模型残差表征观测风险，设计M0、M1、M2模式分配及恢复机制。新任务启动时清除旧命令、积分状态和降级锚点，校验工况配置并建立可信观测；同一任务模式切换时反算积分状态、同步云边状态，结合滞回、驻留时间和恢复确认实现平滑衔接。',
 (2,0,8): '（4）物联网管控系统实现。采用Python开发模型、云端预测、边缘监督和评价模块，围绕样机控制器完成设备接口适配，以MQTT 3.1.1与QoS 1连接云边进程。数据库保存设备、接入点、配方、模型、任务和事件；管理端展示EC、流量压力、双风险、当前模式和命令反馈，支持接入配置检查、作业启动、异常告警与历史追溯。',
 (2,0,9): '（5）分层验证与样机对照。仿真覆盖接入工况、网络状态和观测状态的组合，软件在环核对消息与控制状态，样机试验评价不同压力和管路配置下的启动、稳定作业及短时通信中断。比较云端预测、固定阈值边缘回退和双风险协同控制，并通过启动管理、观测风险及切换机制消融解释性能来源。',
 (2,0,11): '（1）如何依据便携样机不同接入工况的测量数据配置混肥模型，并在转场启动后建立可信控制状态。',
 (2,0,12): '（2）如何联合通信健康与EC观测可信度分配控制权，并实现同一任务内的平滑切换和新任务间的状态隔离。',
 (2,0,13): '（3）如何在共享底层PI、对象工况和执行约束的条件下，评价接入后EC恢复、异常期间跟踪性能与执行器动作代价。',
 (2,0,14): '（4）如何绑定设备、接入点与任务版本，使样机作业与控制服务可追溯。',
 (3,0,2): '技术路线以便携样机的转场接入为起点，依次开展接口与作业流程定义、传感及注肥标定、接入工况动态测量、混肥模型与公共PI、双风险监督和任务初始化、MQTT系统实现，以及配对仿真、软件在环和样机对照。红壤丘陵区的分散生产单元用于组织多接入点需求，设备在完成连接后固定作业；试验围绕接入条件变化和EC控制展开。',
 (3,0,4): '图1  便携式水肥一体机协同管控系统技术路线图',
 (3,0,5): '2. 样机测量、混肥对象与底层控制',
 (3,0,6): '样机试验先完成EC探头、流量和注肥输出标定，再测量不同接入工况下的响应。入口水EC、母液种类与浓度、温度、出口EC、供水压力、流量、管路配置及控制动作同步记录；辨识与验证采用不同的完整启动—稳定作业记录。以增益、时间常数和纯滞后描述混肥动态，保留验证残差作为正常观测风险校准依据。仿真采样与积分步长先经数值检查，样机采样和控制周期依据响应速度及接口能力确定，在正式评价前冻结。',
 (3,0,7): '三种模式共享同一并联离散PI、饱和、抗积分饱和和执行器变化率限制。M0利用时间戳与名义模型补偿观测时延，M1以可信边缘观测闭环，M2以同一任务内最后可信观测校正模型递推。不同接入工况的模型配置由辨识阶段确定，正式对照时各方法使用相同模型与底层参数，比较重点落在监督、启动和恢复机制。',
 (3,0,10): '连续作业以M0为正常模式：通信风险超过阈值且观测可信时进入M1，观测风险升高时进入M2；滞回、最小驻留时间和连续恢复确认管理模式恢复。切换前后通过PI积分反算与云边状态同步衔接控制动作。转场则执行新任务初始化：停止注肥并清除旧命令、积分状态和模型锚点，核对接入点配置与配方，完成通水和可信观测确认后从边缘本地控制启动，再按通信与状态同步条件转入M0。新任务的降级锚点由本次作业建立。',
 (3,0,11): '监督器负责控制位置、反馈来源与恢复条件，任务管理负责跨接入点的配置校验和状态初始化。转场、接入确认、任务启动、模式切换与作业结束均记录事件，包含设备、接入点、源任务、目标任务、模型版本、触发条件和执行反馈，以区分同一作业内的控制切换与不同作业间的状态交接。',
 (3,0,13): '系统按设备适配、边缘监督、云端服务、数据管理和界面进程组织。仿真对象端与样机适配端使用相同消息结构；设备端发布EC和流量压力等状态，云端执行时间戳预测，边缘端负责启动检查、双风险计算和控制执行。MQTT消息携带设备标识、接入点标识、任务标识、序号与源时间戳，接收端执行幂等、陈旧拒绝和跨任务隔离。样机控制器的原有压力液位联锁与人工停止通道保持独立，控制服务经接口适配下发动作。',
 (3,0,14): '数据库按设备—接入点—任务组织模型、配方、参数、观测、动作和事件记录。管理界面以接入检查、作业运行和历史追溯为核心：接入页核对配置与标定状态，运行页展示EC设定值、流量压力、双风险和模式，历史页关联转场、启动、异常与恢复记录。任务结束后形成运行摘要，以同一任务链连接工况、控制过程和评价结果。',
 (3,0,15): '5. 分层试验与评价指标',
 (3,0,16): '仿真采用接入工况、网络状态、观测状态和随机种子配对设计，比较云端时间戳预测、固定阈值边缘回退与双风险协同控制。观测风险消融检验反馈可信度的作用，切换机制消融检验状态衔接，启动管理消融检验跨任务清理与初始化；中断时长扫描评价控制连续性。模型、底层PI和风险参数均在开发阶段确定，独立评价后保持固定。',
 (3,0,17): '样机对照初步设置低、中、高三个供水压力水平与两种管路配置，具体数值限定在设备额定范围和稳定可运行区间内，由预试验确定后冻结。各工况先开展正常启动和稳态作业，再设置可控短时通信中断及重连。每种方法在同一工况下安排不少于3次独立启动—作业重复，按试验批次轮换方法顺序，并统一入口水EC、母液配方、初始混合状态和清洗流程。重复数依据预试验变异与主要终点精度进一步确定。转场试验另比较规范初始化与原有启动流程，记录从接入检查通过至EC持续进入允许带的时间，核对旧任务命令是否被拒绝。',
 (3,0,18): '主要控制指标为EC超出允许带的累计时间Tout，辅助指标为IAE、最大误差、稳定时间、执行器总变差TV、切换次数及模式占用率。转场作业增加安装调试时间、接入后EC恢复时间、启动成功率和任务隔离检查；设备质量与尺寸用于描述便携作业条件。EC允许带、持续达标判据和计时起点在正式试验前统一，恢复失败按失败记录并纳入汇总。以完整试验运行作为重复单位，配对报告差值、区间与尾部工况。仿真、软件在环与样机结果分层分析，通信和传感异常的广泛组合由仿真与软件在环承担，样机试验验证代表性接入和恢复过程。',
 (3,1,1): '（1）面向便携水肥装备的转场接入协同管理。拟将接入点配置、混肥模型、任务初始化与控制模式管理连接起来，区分跨任务状态重建和同一任务内的无扰切换，以接入后EC恢复、旧命令隔离及独立启动重复检验作业衔接机制。',
 (3,1,2): '（2）通信与观测双风险驱动的控制权分配。拟联合命令时效性与本地EC观测可信度，在云端预测、边缘控制和模型锚定降级之间选择控制位置与反馈来源；通过观测风险和切换机制消融，检验各环节对异常期间EC调节与执行器动作的作用。',
 (3,1,3): '（3）样机数据支撑的接入工况表征与分层评价。拟以供水压力、流量和管路配置下的实测混肥响应确定模型与工况，再将配对仿真、MQTT软件在环和样机对照连接到同一任务规范，形成便携作业条件、控制机制与系统实现相互对应的验证链。',
 (3,1,4): '便携式样机是上述研究的装备载体，丘陵分散地块的轮换接入是应用特色；拟研究的创新集中于转场状态管理、双风险控制权分配及其样机验证。',
 (3,2,2): '本课题采用便携式水肥一体机作为实体试验载体，围绕供水、注肥、EC检测、设备状态和控制接口建立测量与接入方案。样机台架用于标定、动态辨识和不同接入工况下的对照试验；本地计算环境用于对象仿真、异常注入、统计评价与MQTT软件在环，构成从模型到控制软件再到实体运行的验证流程。',
 (3,2,5): '（1）接入条件变化会改变混肥动态。先开展额定范围内的压力、流量与管路响应试验，分别确定辨识和验证数据；为接入点绑定工况配置与模型版本，以正常残差校准观测风险，并在模型失配场景中检查误切换。',
 (3,2,7): '（3）转场初始化与模式切换需要不同的状态处理。新任务清除旧命令和旧锚点，重建传感与模型状态；同一任务内采用PI积分反算、云边状态同步和变化率限制衔接输出。分别用跨任务命令拒绝检查、启动恢复试验与切换TV评价其作用。',
 (3,2,10): '（6）样机初始状态和配方残留会影响重复试验。规定通水、清洗、母液配制及初始EC确认流程，记录温度、压力和流量；预试验确定测量周期、允许带和稳定判据后冻结。样机试验保留独立参考EC、压力液位联锁与人工停止通道，在线异常试验按预先规定的可控范围实施。',
}

def set_text(p,text):
    base=deepcopy(p.runs[0]._element.rPr) if p.runs and p.runs[0]._element.rPr is not None else None
    p.clear()
    for part in re.split(r'(\[[0-9,-]+\])',text):
        if not part: continue
        r=p.add_run(part)
        if base is not None: r._element.insert(0,deepcopy(base))
        r.font.superscript=bool(re.fullmatch(r'\[[0-9,-]+\]',part))
        if r.font.superscript: r.font.size=Pt(9)

def make_route(path):
    # Reuse the established vector/layout logic; redraw labels to match this revision.
    src=(ROOT/'scripts/redesign_proposal_technical_route.py').read_text()
    changes={
      '控制对象与边界\n混肥过程  EC允许带  执行器约束':'便携作业与接入工况\n供水压力  流量  管路  初始EC',
      '异常与信息流\n时延丢包中断  偏置漂移尖峰':'任务交接与异常\n转场启动  时延中断  观测异常',
      '评价问题\n跟踪性能  动作代价  切换平滑性':'评价问题\n接入后恢复  EC跟踪  动作平滑',
      '物联网管控架构\n感知执行—边缘—MQTT—云端—管理':'样机测量与工况配置\nEC流量压力标定  辨识与独立验证',
      '消息与任务模型\n时间戳  序号  幂等  版本与事件审计':'接入点与任务管理\n配置核对  旧状态清理  可信启动',
      '滞回  驻留时间  恢复确认  状态同步':'滞回  恢复确认  同任务状态同步',
      '对象参数 × 网络异常 × 传感异常 × 随机种子':'接入工况 × 网络异常 × 观测异常',
      '云端基线  固定回退  风险消融  中断时长扫描':'云端基线  固定回退  启动与风险消融',
      'MQTT软件在环\n消息异常  模式序列  状态同步  进程恢复':'软件在环与样机对照\n启动恢复  工况配对  消息与状态核对',
      '性能与边界评价：Tout  IAE  最大误差  稳定时间  TV  切换次数  模式占用率  尾部工况':'性能评价：接入后EC恢复  Tout  IAE  TV  切换次数  启动成功率  任务隔离',
      '物联网边云协同管控原型\n接口与任务审计规范':'便携样机协同管控系统\n接入配置与作业追溯',
      '双风险三模式监督方法\n参数与适用条件':'转场初始化与三模式监督\n状态衔接与恢复机制',
      '配对评价与软件在环报告\n失效边界与扩展接口':'样机测量与分层对照报告\n接入条件与性能评价',
    }
    for a,b in changes.items(): src=src.replace(a.replace('\n','\\n'),b.replace('\n','\\n'))
    ns={'__name__':'portable_route_helper'}
    exec(compile(src,'portable_route_helper','exec'),ns)
    ns['make_technical_route'](path,'红壤丘陵区分散地块便携水肥装备的转场接入与EC协同控制')

def main():
    ASSETS.mkdir(parents=True,exist_ok=True)
    d=Document(SOURCE)
    for (ti,ri,pi),text in REPLACEMENTS.items(): set_text(d.tables[ti].rows[ri].cells[0].paragraphs[pi],text)
    p=d.tables[1].cell(0,0).paragraphs[2]
    set_text(p,p.text.replace('自主水肥管理系统和基于物联网的微服务框架，将水源管理、数据服务与水肥执行连接起来','自主水肥管理系统和物联网微服务框架，将水源管理、数据服务与水肥执行相连接'))
    p=d.tables[0].cell(0,0).paragraphs[3]
    set_text(p,REPLACEMENTS[(0,0,3)].replace('由此，将转场后的任务初始化与作业中的模式切换连接起来，是本课题系统设计与控制方法的共同切入点。','本课题据此统一设计转场任务初始化与作业内模式切换。'))
    for ti,ri,pi in REPLACEMENTS:
        d.tables[ti].rows[ri].cells[0].paragraphs[pi].paragraph_format.widow_control=True
    schedule={
      3:('完成样机接口与传感注肥标定，测量压力和管路工况下的混肥响应，开展独立模型验证。','接入规格、标定记录与混肥模型'),
      4:('实现双风险三模式监督、转场任务初始化和无扰切换，冻结开发配置。','协同控制策略与启动恢复流程'),
      5:('开展接入工况与异常组合配对仿真、基线对照和启动及风险机制消融。','配对评价和机制分析'),
      6:('完成MQTT软件在环、跨任务隔离、消息与状态一致性核对，适配样机接口。','软件在环报告与样机适配模块'),
      7:('完成接入配置、作业运行与历史追溯界面，开展样机预试验并冻结对照协议。','管控系统与样机试验协议'),
      8:('开展不同供水压力和管路配置下的样机对照、转场启动及短时通信中断恢复试验。','样机对照数据与综合评价'),
    }
    for ri,(a,b) in schedule.items():
        set_text(d.tables[4].rows[ri].cells[1].paragraphs[0],a)
        set_text(d.tables[4].rows[ri].cells[2].paragraphs[0],b)
    outline=d.tables[4].rows[16].cells[0].paragraphs
    titles=['九、论文详细提纲','第1章 绪论：红壤丘陵区分散地块应用、便携水肥装备进展、研究问题与技术路线。','第2章 便携式水肥一体机管控系统总体设计：装备接口、转场接入、分层架构与任务流程。','第3章 样机动态测量与混肥EC建模：传感及注肥标定、接入工况、辨识验证与公共PI。','第4章 风险感知协同控制方法：双风险、三模式监督、转场初始化与无扰切换。','第5章 仿真与软件在环验证：配对场景、基线、机制消融、消息一致性和任务隔离。','第6章 系统实现与便携样机试验：设备适配、管理界面、压力和管路工况对照、启动与恢复。','第7章 综合分析：接入工况、EC调节、执行器动作、共同漂移及模型适用范围。','第8章 结论与展望：系统设计、控制方法、分层验证与装备应用。']
    for p,t in zip(outline,titles): set_text(p,t)
    refs=d.tables[4].rows[15].cells[0]
    r=refs.add_paragraph()
    template=refs.paragraphs[-2]
    if template._p.pPr is not None: r._p.insert(0,deepcopy(template._p.pPr))
    set_text(r,'[37] 左志宇, 魏丰瑞, 毛罕平, 等. 多通道移动式果园灌溉施肥机设计与试验[J]. 农机化研究, 2023, 45(4): 95-100. DOI: https://doi.org/10.13427/j.cnki.njyi.2023.04.008')
    for run in r.runs:
        run.font.superscript=False; run.font.size=Pt(9); run.font.name='Times New Roman'
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    refs.paragraphs[0].paragraph_format.keep_with_next=True
    for p in refs.paragraphs[1:]:
        p.paragraph_format.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    # Keep the school approval forms starting on a complete page.
    d.tables[5].rows[0].cells[0].paragraphs[0].paragraph_format.page_break_before=True
    d.tables[5].rows[2].cells[0].paragraphs[0].paragraph_format.page_break_before=True
    from docx.oxml import OxmlElement
    for row in d.tables[5].rows:
        pr=row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None: pr.append(OxmlElement('w:cantSplit'))
    from docx.text.paragraph import Paragraph
    for e in list(d.element.body):
        if e.tag==qn('w:p') and e.getprevious() is not None and e.getnext() is not None:
            if e.getprevious().tag==qn('w:tbl') and e.getnext().tag==qn('w:tbl'):
                p=Paragraph(e,d)
                if not p.text and not e.xpath('.//w:drawing|.//w:pict'):
                    p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=Pt(1)
                    p.paragraph_format.keep_with_next=False
                    p.paragraph_format.page_break_before=False
                    for r in p.runs: r.font.size=Pt(1)
    # Non-breaking hyphens keep numeric citation ranges on one line.
    for e in d.element.xpath('.//w:t'):
        if e.text and re.fullmatch(r'\[[0-9,-]+\]',e.text): e.text=e.text.replace('-','\u2011')
    d.save(OUTPUT)
    route=ASSETS/'technical_route_portable.png'; make_route(route)
    # Find the route by its actual drawing relationship rather than media filename assumptions.
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(OUTPUT) as z:
        tree=etree.fromstring(z.read('word/document.xml'))
    pics=tree.xpath('//*[local-name()="blip" or local-name()="imagedata"]')
    embed=pics[-1].get(qn('r:embed')) or pics[-1].get(qn('r:id'))
    target='word/'+str(d.part.rels[embed].target_ref)
    rebuild_with_media(OUTPUT,{target:route})
    archive=ROOT/'output/proposal_revision/开题报告_便携式样机与协同控制修订版.docx'
    shutil.copy2(OUTPUT,archive)
    checked=Document(OUTPUT); body=[]; seen=set()
    for ti,t in enumerate(checked.tables):
        for ri,row in enumerate(t.rows):
            for c in row.cells:
                if c._tc in seen: continue
                seen.add(c._tc)
                if ti==4 and ri==15: continue
                body.extend(p.text.replace('\u2011','-') for p in c.paragraphs)
    order=[]; groups=[]
    for text in body:
        for m in re.finditer(r'\[([0-9,-]+)\]',text):
            vals=[]
            for s in m.group(1).split(','):
                if '-' in s:
                    a,b=map(int,s.split('-')); vals.extend(range(a,b+1))
                else: vals.append(int(s))
            assert len(vals)<=4; groups.append(vals)
            for n in vals:
                if n not in order: order.append(n)
    assert order==list(range(1,38)),order
    joined='\n'.join(body)
    assert 'RA-ECSC' not in joined
    assert not any(s in joined for s in ['项目未开始','拟参加','已取得','已完成算法','已经实现','1500组'])
    with ZipFile(OUTPUT) as z:
        assert z.testzip() is None
        assert len([n for n in z.namelist() if n.startswith('word/media/')])==2
    data={'output':str(OUTPUT),'modified_paragraphs':len(REPLACEMENTS),'references':37,'books':11,'max_citations_per_group':max(map(len,groups)),'first_appearance':order,'images':2,'new_reference_verified':'https://www.aeeisp.com/njhyj/article/2023/4','prototype_results_assumed':False,'mainline':'便携样机转场接入管理与混肥EC双风险三模式协同控制'}
    (ASSETS/'audit.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    with (ROOT/'data/csv/proposal_references_t060.csv').open(encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f); fields=reader.fieldnames; rows=list(reader)
    rows.append(dict(zip(fields,[37,'J',2023,'左志宇, 魏丰瑞, 毛罕平, 等. 多通道移动式果园灌溉施肥机设计与试验[J]. 农机化研究, 2023, 45(4): 95-100.','https://doi.org/10.13427/j.cnki.njyi.2023.04.008','no','期刊官方文章页与2023年第4期目录交叉核验','verified'])))
    with (ROOT/'data/csv/proposal_references_t062.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields,lineterminator='\n'); writer.writeheader(); writer.writerows(rows)
    (ASSETS/'writing_review.md').write_text('''# T062 修订依据与论证核对

主线：面向便携样机在分散接入点的轮换作业，以实测混肥动态配置对象，以任务初始化隔离转场旧状态，以双风险三模式监督管理连续作业，并经配对仿真、软件在环和样机试验验证。

## 术语

- 便携式水肥一体机：转场后连接供水与灌溉管网并固定作业的装备载体；不预设行走、牵引或电池结构。
- 转场任务初始化：新任务核对配置、清除旧命令与旧锚点并重建可信状态。
- 无扰切换：同一任务内通过积分反算与云边状态同步衔接控制动作。
- M0/M1/M2：云端预测、边缘控制、模型锚定降级；底层共用PI。
- EC：混合肥液电导率反馈；不等同于具体养分浓度或作物肥效。

## 依据与待验证主张

- 可接触样机：用户确认；额定范围、传感配置、控制器接口及可修改方式尚待现场核对。
- 便携/移动装备已有研究：期刊官方文章页和同期目录支持；引用[37]，不将移动性本身表述为新发明。
- 接入条件对模型和恢复的影响：研究假设，拟由压力与管路工况下的样机测量验证。
- 任务初始化价值：拟由启动流程对照、旧命令拒绝及启动管理消融验证。
- 双风险监督价值：拟由共享PI条件下的对照、观测风险消融及切换机制消融验证。
- 系统一致性：拟由软件在环逐点核对和样机事件审计验证。

正式稿仅陈述拟研究内容、方法及验证设计，不填入样机参数、既有试验结果或区域节水增产主张。不把任务设计中的重复数写成已完成样本数。
''',encoding='utf-8')
    print(json.dumps(data,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
