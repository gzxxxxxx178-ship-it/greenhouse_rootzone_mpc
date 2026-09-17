from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw

from revise_proposal_from_current_evidence import (
    arrow,
    box,
    font,
    normalize_document_fonts,
    rebuild_with_media,
    set_text,
    wrapped,
)


SOURCE = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于MPC的智能水肥一体化系统设计研究_设计稿.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于风险感知边云协同控制的智能水肥一体化系统设计与实现.docx"
)


def set_paragraphs(paragraphs, updates: dict[int, str]) -> None:
    for index, value in updates.items():
        set_text(paragraphs[index], value)


def make_architecture(path: Path) -> None:
    image = Image.new("RGB", (2362, 1921), "white")
    draw = ImageDraw.Draw(image)
    box(draw, (250, 50, 2110, 280), "感知与执行层\n双EC传感器  流量压力液位  注肥泵  混肥罐  阀门", title=True)
    box(draw, (420, 390, 1940, 640), "MQTT通信与消息语义\nQoS 1  序号与时间戳  幂等处理  状态同步", title=True)
    box(draw, (180, 790, 940, 1190), "边缘控制层\n数据预处理与时间戳\n本地PI控制  风险计算\n状态机  联锁与降级执行")
    box(draw, (1420, 790, 2180, 1190), "云端服务层\nEC预测补偿  模型与参数管理\n任务编排  数据存储\n告警审计与管理界面")
    box(draw, (250, 1480, 2110, 1780), "风险感知边云协同控制\nM0云端预测  M1边缘控制  M2模型锚定降级\n滞回  驻留时间  无扰切换", title=True)
    arrow(draw, (1180, 280), (1180, 390))
    arrow(draw, (850, 640), (620, 790)); arrow(draw, (1510, 640), (1800, 790))
    arrow(draw, (620, 1190), (850, 1480)); arrow(draw, (1800, 1190), (1510, 1480))
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))


def make_supervisor(path: Path) -> None:
    image = Image.new("RGB", (1689, 1338), "white")
    draw = ImageDraw.Draw(image)
    box(draw, (110, 70, 720, 290), "通信风险 R_net\n命令年龄  时延EWMA\n窗口丢包率", title=True)
    box(draw, (970, 70, 1580, 290), "观测风险 R_obs\n双传感器差异\n融合值与模型残差", title=True)
    box(draw, (420, 430, 1260, 690), "风险感知监督状态机\n进入阈值与恢复阈值分离\n最小驻留与恢复确认", title=True)
    box(draw, (70, 850, 520, 1110), "M0 云端预测\n时间戳补偿\n正常网络优先")
    box(draw, (620, 850, 1070, 1110), "M1 边缘控制\n本地可信观测\n通信异常接管")
    box(draw, (1170, 850, 1620, 1110), "M2 锚定降级\n最后可信观测\n模型预测反馈")
    box(draw, (420, 1160, 1260, 1320), "状态同步与无扰切换\n统一执行器变化率约束", title=True)
    arrow(draw, (420, 290), (650, 430)); arrow(draw, (1270, 290), (1030, 430))
    arrow(draw, (650, 690), (295, 850)); arrow(draw, (840, 690), (845, 850)); arrow(draw, (1030, 690), (1395, 850))
    arrow(draw, (295, 1110), (620, 1160)); arrow(draw, (845, 1110), (845, 1160)); arrow(draw, (1395, 1110), (1070, 1160))
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))


def make_validation(path: Path) -> None:
    image = Image.new("RGB", (817, 640), "white")
    draw = ImageDraw.Draw(image)
    normal = font(18); bold = font(20, bold=True)

    def small(x1, y1, x2, y2, text, fill="#EEF4F8", strong=False):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=8, fill=fill, outline="#6B7C8C", width=2)
        selected = bold if strong else normal
        lines = wrapped(draw, text, selected, x2 - x1 - 16)
        y = y1 + (y2 - y1 - 24 * len(lines)) / 2
        for line in lines:
            b = draw.textbbox((0, 0), line, font=selected)
            draw.text((x1 + (x2 - x1 - (b[2] - b[0])) / 2, y), line, font=selected, fill="#1F2933")
            y += 24

    small(80, 20, 737, 82, "对象模型与公共底层PI冻结", "#E6F0F7", True)
    small(35, 125, 245, 210, "5种对象\n参数变化与切换")
    small(303, 125, 513, 210, "6种网络\n时延丢包与中断")
    small(571, 125, 781, 210, "5种观测\n偏置漂移与尖峰")
    small(80, 255, 737, 330, "开发集选择参数  独立评价集冻结验证", "#E6F0F7", True)
    small(25, 385, 245, 470, "总体组合评价\n1500配对场景")
    small(298, 385, 518, 470, "时长扫描与消融\n适用条件和机制")
    small(571, 385, 791, 470, "MQTT软件在环\n跨进程一致性")
    small(80, 545, 737, 625, "Tout  IAE  TV  稳定时间  切换次数  模式占用率", "#FFF3DF", True)
    for x in (140, 408, 676): arrow(draw, (408, 82), (x, 125), width=4)
    for x in (140, 408, 676): arrow(draw, (x, 210), (408, 255), width=4)
    for x in (135, 408, 681): arrow(draw, (408, 330), (x, 385), width=4)
    for x in (135, 408, 681): arrow(draw, (x, 470), (408, 545), width=4)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))


REFERENCES = [
    "[1] WU Y, LI L, LI S, et al. Optimal control algorithm of fertigation system in greenhouse based on EC model[J]. International Journal of Agricultural and Biological Engineering, 2019, 12(3): 118-125. DOI:10.25165/j.ijabe.20191203.4680.",
    "[2] XU Y, JIN Y, SUN Z, XUE X. PSO-based system identification and fuzzy-PID control for EC real-time regulation in fertilizer mixing system[J]. Agronomy, 2025, 15(5): 1259. DOI:10.3390/agronomy15051259.",
    "[3] YUAN W, ZHANG Z, DAI J, et al. Change-aware online identification and adaptive PI control for robust electrical conductivity regulation in fertigation mixing systems[J]. Smart Agricultural Technology, 2026, 14: 102110. DOI:10.1016/j.atech.2026.102110.",
    "[4] GONZÁLEZ PEREA R, MORENO M A, ORTEGA J F, et al. Dynamic simulation tool of fertigation in drip irrigation subunits[J]. Computers and Electronics in Agriculture, 2020, 173: 105434. DOI:10.1016/j.compag.2020.105434.",
    "[5] SUN G, LI X, WANG X, et al. Design and testing of a nutrient mixing machine for greenhouse fertigation[J]. Engineering in Agriculture, Environment and Food, 2015, 8(2): 114-121. DOI:10.1016/j.eaef.2014.12.001.",
    "[6] 熊钦, 肖丽萍, 蔡金平, 等. 基于物联网的果园药水肥一体化控制系统设计与实现[J]. 中国农机化学报, 2023, 44(3): 73-81. DOI:10.13733/j.jcam.issn.2095-5553.2023.03.011.",
    "[7] SHI W, CAO J, ZHANG Q, et al. Edge computing: Vision and challenges[J]. IEEE Internet of Things Journal, 2016, 3(5): 637-646. DOI:10.1109/JIOT.2016.2579198.",
    "[8] WANG Z, HUANG J, CHEN C, FUKUSHIMA S. Design of prediction-based controller for networked control systems with packet dropouts and time-delay[J]. Mathematical Problems in Engineering, 2022: 9437955. DOI:10.1155/2022/9437955.",
    "[9] HESPANHA J P, NAGHSHTABRIZI P, XU Y. A survey of recent results in networked control systems[J]. Proceedings of the IEEE, 2007, 95(1): 138-162. DOI:10.1109/JPROC.2006.887288.",
    "[10] ZHANG L, GAO H, KAYNAK O. Network-induced constraints in networked control systems: A survey[J]. IEEE Transactions on Industrial Informatics, 2013, 9(1): 403-416. DOI:10.1109/TII.2012.2219540.",
    "[11] ISERMANN R. Model-based fault-detection and diagnosis: Status and applications[J]. Annual Reviews in Control, 2005, 29(1): 71-85. DOI:10.1016/j.arcontrol.2004.12.002.",
    "[12] FENG J, TURKSOY K, SAMADI S, et al. Hybrid online sensor error detection and functional redundancy for systems with time-varying parameters[J]. Journal of Process Control, 2017, 60: 115-127. DOI:10.1016/j.jprocont.2017.04.004.",
    "[13] MA Y, LU C, SINOPOLI B, et al. Exploring edge computing for multitier industrial control[J]. IEEE Transactions on Computer-Aided Design of Integrated Circuits and Systems, 2020, 39(11): 3506-3518. DOI:10.1109/TCAD.2020.3012648.",
    "[14] ZAREI J, MASOUDI E, RAZAVI-FAR R, et al. Fault-tolerant control design for unreliable networked control systems via constrained model predictive control[J]. ISA Transactions, 2023, 134: 171-182. DOI:10.1016/j.isatra.2022.08.019.",
    "[15] YETENDJE A, DE DONÁ J A, SERON M M. Multisensor fusion fault tolerant control[J]. Automatica, 2011, 47(7): 1461-1466. DOI:10.1016/j.automatica.2011.02.024.",
    "[16] SERON M M, DE DONÁ J A, OLARU S. Fault tolerant control allowing sensor healthy-to-faulty and faulty-to-healthy transitions[J]. IEEE Transactions on Automatic Control, 2012, 57(7): 1657-1669. DOI:10.1109/TAC.2011.2178716.",
    "[17] TRAN T C, JUNG J C. Development of anti-windup PI control and bumpless control transfer methodology for feedwater control system[J]. Annals of Nuclear Energy, 2019, 131: 233-241. DOI:10.1016/j.anucene.2019.03.030.",
    "[18] ZACCARIAN L, TEEL A R. The L2 bumpless transfer problem for linear plants: Its definition and solution[J]. Automatica, 2005, 41(7): 1273-1280. DOI:10.1016/j.automatica.2005.02.003.",
    "[19] OASIS. MQTT Version 3.1.1[S/OL]. 2014. https://docs.oasis-open.org/mqtt/mqtt/v3.1.1/mqtt-v3.1.1.html.",
    "[20] KWON W H, CHOI S G, KIM K B. Network-based software-in-the-loop simulation for real-time control system[J]. IFAC Proceedings Volumes, 1999, 32(2): 6047-6052. DOI:10.1016/S1474-6670(17)57032-9.",
    "[21] DARDER M, VALERA A, NIETO E, et al. Multisensor device based on case-based reasoning for monitoring nutrient solutions in fertigation[J]. Sensors and Actuators B: Chemical, 2009, 135(2): 530-536. DOI:10.1016/j.snb.2008.09.034.",
    "[22] WANG H, ZHAO J, ZHANG L, et al. Application of disturbance observer-based fast terminal sliding mode control for remote electrical conductivity control of fertigation systems[J]. Agriculture, 2024, 14(2): 168. DOI:10.3390/agriculture14020168.",
    "[23] ZOU X, LIU W, HUO Z, et al. Current status and prospects of research on sensor fault diagnosis of agricultural Internet of Things[J]. Sensors, 2023, 23(5): 2528. DOI:10.3390/s23052528.",
    "[24] IMBERNÓN-MULERO A, MAESTRE-VALERO J F, MARTÍNEZ-ALVAREZ V, et al. Evaluation of an autonomous smart system for optimal management of fertigation with variable sources of irrigation water[J]. Frontiers in Plant Science, 2023, 14: 1149956. DOI:10.3389/fpls.2023.1149956.",
    "[25] ZHANG H, HE L, GIOIA F D, et al. LoRaWAN based Internet of Things system for precision irrigation in plasticulture fresh-market tomato[J]. Smart Agricultural Technology, 2022, 2: 100053. DOI:10.1016/j.atech.2022.100053.",
    "[26] ADAMO T, CAIVANO D, COLIZZI L, et al. Optimization of irrigation and fertigation in smart agriculture: An IoT-based micro-services framework[J]. Smart Agricultural Technology, 2025, 11: 100885. DOI:10.1016/j.atech.2025.100885.",
    "[27] KALYANI Y, BERMEO N V, COLLIER R. Digital twin deployment for smart agriculture in Cloud-Fog-Edge infrastructure[J]. International Journal of Parallel, Emergent and Distributed Systems, 2023, 38(6): 461-476. DOI:10.1080/17445760.2023.2235653.",
    "[28] SHEKARIAN S M, AMINIAN M, FALLAH A M, et al. AI-powered sensor fault detection for cost-effective smart greenhouses[J]. Computers and Electronics in Agriculture, 2024, 224: 109198. DOI:10.1016/j.compag.2024.109198.",
]


def main() -> None:
    doc = Document(SOURCE)
    title = "论文题目：基于风险感知边云协同控制的智能水肥一体化系统设计与实现"
    set_text(doc.paragraphs[6], title)
    for run in doc.paragraphs[6].runs:
        run.font.size = Pt(17)

    set_paragraphs(doc.tables[0].rows[0].cells[0].paragraphs, {
        2: "水肥一体机通过注肥泵将母液按比例注入灌溉水，并以电导率（electrical conductivity，EC）表征混合肥液离子浓度。混肥过程具有惯性、纯滞后和工况变化，EC闭环既要保持浓度跟踪精度，也要限制注肥执行器的频繁动作。随着控制任务由单机本地运行转向云端服务与边缘节点协同，测量、命令和状态需要经过网络传输，时延、丢包、重复和通信中断开始直接影响闭环性能[1-10]。",
        3: "常规云端预测补偿能够减弱轻中度网络波动，但持续中断时需要边缘端接管控制。固定超时回退只回答通信是否正常，无法判断本地EC观测是否可信；当通信异常与传感器偏置、漂移或尖峰并发时，直接把控制权交给边缘PI可能放大错误反馈。水肥一体化系统因此需要同时评估通信健康和观测可信度，并在不同控制模式之间完成平稳、可追溯的控制权转移。",
        4: "2. 选题定位",
        5: "本课题拟设计并实现基于风险感知边云协同控制的智能水肥一体化系统。系统以双EC传感、注肥执行、边缘节点、MQTT通信、云端控制服务和管理平台为主体，以风险感知自适应边云监督控制（RA-ECSC）为核心方法，在云端预测模式、边缘控制模式和模型锚定降级模式之间分配控制权。底层控制器采用统一PI、饱和、抗积分饱和和动作变化率约束，监督器负责风险计算、模式切换、状态同步和异常回退。",
        6: "研究从系统、算法和实现三条线同步展开：系统层定义传感、通信、控制、执行与审计接口；算法层研究通信风险、观测可信风险和三模式状态机；实现层完成跨进程MQTT通信、任务与版本记录、控制服务和管理界面。验证采用对象参数变化、网络异常、传感异常、机制消融、通信中断时长扫描及软件在环组合，条件具备时再接入EC探头和注肥执行器开展台架复核。",
        7: "3. 研究意义",
        8: "本研究把水肥一体机EC控制从单一控制律设计推进到控制权管理问题。通信风险决定云端命令是否仍具有时效性，观测可信度决定本地反馈是否适合继续参与闭环，两类信息共同决定控制位置和降级方式。形成的系统不仅能够调节EC，还能记录风险、模式、同步、命令和反馈，为异常工况下的控制策略评价、软件复现和装备部署提供统一载体。",
    })

    set_paragraphs(doc.tables[1].rows[0].cells[0].paragraphs, {
        1: "1. 水肥装备与EC闭环控制",
        2: "水肥一体化装备研究已覆盖多通道混配、营养液EC与pH监测、自动注肥及物联网管理[1-6]。针对EC动态，已有工作采用模型控制、系统辨识、模糊PID、在线辨识与自适应PI改善浓度跟踪[1-5]。这些研究说明EC能够作为混肥过程的核心反馈变量，也表明对象增益、惯性、时滞和母液工况变化会直接影响控制器参数。现有成果多聚焦对象与底层控制律，对网络化部署后的控制权转移讨论相对不足。",
        3: "2. 网络化控制与边云协同",
        4: "边缘计算为现场快速控制和云端模型服务之间的分层协同提供了基础[7,13]。网络化控制研究指出，时延、丢包和数据乱序会改变反馈时序，预测补偿、状态估计和容错控制可提高闭环连续性[8-10,14]。对水肥一体机而言，云端具有集中计算和模型管理优势，边缘端具有低时延与局部自治优势；二者的关键不在于固定选择某一位置，而在于根据运行风险动态配置控制权。",
        5: "通信恢复并不意味着可以立即回到云端控制。陈旧命令、积分状态不一致和连续模式切换可能造成执行器突变。滞回、最小驻留时间、恢复确认、陈旧消息清除和控制器状态同步是实现无扰转移的重要机制[17-18]，需要与具体消息时序和执行器约束共同设计。",
        6: "3. 传感异常与可信监督",
        7: "EC传感器可能出现固定偏置、线性漂移、尖峰和共同漂移。双传感器差异能够识别通道不一致，模型预测残差能够补充同向异常信息，但残差同时包含对象失配、建模误差与测量异常[11-12,15-16,21,23,28]。因此，观测风险更适合用于控制权决策和降级管理，而不能直接等同于故障分类结果。异构参考测量、在线校准和专用诊断可作为共同漂移条件下的扩展。",
        8: "4. 研究述评与趋势",
        9: "现有研究分别解决了水肥装备、EC底层控制、网络补偿、边缘计算和传感故障诊断问题，但缺少面向水肥一体机的统一控制权监督与可复现系统实现。本课题据此围绕“双风险评估—三模式分配—无扰切换—MQTT实现—分层验证”建立研究链，使算法贡献能够落实到可运行的智能水肥一体化系统。",
    })

    set_paragraphs(doc.tables[2].rows[0].cells[0].paragraphs, {
        2: "设计并实现基于风险感知边云协同控制的智能水肥一体化系统，完成双EC与设备状态采集、混肥对象建模、云端预测控制、边缘PI控制、模型锚定降级控制、风险监督、无扰切换、MQTT通信和任务审计。通过冻结的仿真与软件在环协议，评价通信和传感异常并发条件下的EC调节性能、执行器动作和平滑切换能力，形成可复现的软件系统与验证证据。",
        3: "系统总体架构如图1所示。双EC、流量、压力、液位和执行器状态在边缘侧完成时间戳、质量标识和预处理；云端承担预测补偿、模型版本和任务管理；监督器根据通信风险与观测风险选择M0、M1或M2。MQTT负责跨进程消息交换，数据库保存观测、风险、模式、命令、反馈和版本，管理界面展示实时状态、告警与历史任务。",
        4: "图1  风险感知边云协同智能水肥一体化系统总体架构",
        6: "（1）智能水肥一体化系统总体设计。建立感知层、边缘控制层、通信层、云端服务层、执行层和管理层，定义EC、设备状态、控制命令、模式事件和反馈消息的数据结构，形成系统需求、模块接口与任务流程。",
        7: "（2）混肥对象辨识与底层EC控制。以混肥罐、注肥泵和管路为控制对象，建立一阶惯性纯滞后模型及工况变化对象族；设计统一PI、饱和、抗积分饱和和动作变化率限制，为云端与边缘模式提供一致的底层控制基础。",
        8: "（3）风险感知边云监督控制。利用命令年龄、时延EWMA和窗口丢包率构造通信风险，利用双传感器差异和模型残差构造观测可信风险；设计M0云端预测、M1边缘控制和M2模型锚定降级三模式状态机，并加入滞回、驻留时间、恢复确认和状态同步。",
        9: "（4）MQTT与平台软件实现。采用Python实现对象仿真、边缘监督、云端控制与评价服务，MQTT 3.1.1实现发布订阅和QoS 1消息传输，SQLite或关系数据库保存任务与事件；管理端提供EC曲线、风险分量、控制模式、命令反馈、告警和结果查询。",
        10: "（5）分层验证与适用范围分析。设置对象参数变化、网络时延丢包中断、传感偏置漂移尖峰与共同漂移，开展主要对照、固定阈值回退、机制消融、通信中断时长扫描和MQTT软件在环；进一步分析模型失配、共同漂移和执行器约束对控制收益的影响。",
        12: "（1）如何把通信状态和本地EC观测可信度转化为可在线计算、可解释的控制权分配依据。",
        13: "（2）如何在云端预测、边缘控制和降级控制之间实现稳定切换，并避免积分状态不一致和陈旧命令引起的执行器冲击。",
        14: "（3）如何在相同对象、网络、传感扰动和执行器约束下评价监督策略的跟踪收益、尾部风险和动作代价。",
        15: "（4）如何通过MQTT消息语义、任务版本和事件审计保证参考算法与跨进程实现的一致性。",
    })

    method = doc.tables[3].rows
    set_paragraphs(method[0].cells[0].paragraphs, {
        2: "研究沿“需求与对象定义—混肥模型辨识—公共底层PI—通信与观测风险—三模式监督状态机—无扰切换—配对仿真—MQTT软件在环—系统集成”展开。底层控制器和执行器约束先统一，监督方法只改变控制权与反馈来源，使配对差集中反映风险信息和模式分配逻辑的作用。",
        3: "2. 混肥对象与底层控制",
        4: "以注肥控制量为输入、混肥EC偏差为输出建立一阶惯性纯滞后模型，参数包括稳态增益、时间常数和纯滞后。名义模型用于控制器和预测器，对象评价侧设置低增益慢响应、高增益快响应、时滞主导和运行中参数切换等变体，以检验模型失配。EC按1 s采样，控制命令按2 s更新，数值步长设置为0.2 s。",
        5: "所有模式共享同一并联离散PI、饱和、抗积分饱和和执行器变化率限制。云端依据带时间戳的观测和名义模型预测当前EC；边缘端使用本地可信观测闭环；模型锚定降级模式以最后可信观测校正名义模型预测。公共底层控制使研究重点保持在监督器，而不是不同控制律参数之间的差异。",
        6: "3. 风险监督与控制权分配",
        7: "通信风险由命令年龄、已接收命令时延EWMA和最近窗口序号缺口构成，持续中断会直接推高命令年龄分量。观测风险由双EC通道差和融合观测相对名义模型的残差构成，残差采用EWMA抑制单点噪声。各分量归一化至0至1，风险阈值由独立开发数据冻结。",
        8: "状态机从M0启动：观测可信而通信风险升高时进入M1；观测风险升高时优先进入M2。恢复阈值低于进入阈值，并设置最小驻留与连续恢复确认。进入M1或M2时根据目标连续控制量反算PI积分状态，返回M0时同步云端状态、清除陈旧待执行命令并重新开始命令年龄计算。",
        9: "图2给出了监督器的控制权分配流程。监督器不直接构造新的底层控制律，而是决定当前使用的反馈来源、控制位置和降级机制；每次模式转换同时记录触发风险、来源模式、目标模式、同步状态和最终动作。",
        10: "图2  风险感知三模式监督与无扰切换流程",
        11: "4. 系统软件实现",
        12: "系统采用模块化进程实现：对象端发布EC与设备状态，故障代理注入时延、丢包、重复和传感异常，云端进程执行时间戳预测补偿，边缘进程计算风险并运行监督状态机。MQTT 3.1.1采用QoS 1和非保留消息；接收端使用序号、时间戳和任务标识处理重复、陈旧与跨任务消息。",
        13: "数据库以任务为主线保存对象版本、控制器参数、风险阈值、随机种子、观测、命令、反馈、模式事件和评价结果。管理界面围绕实时EC、设定值、R_net、R_obs、当前模式、执行器输出、告警和历史任务组织，使控制过程、异常原因和实验结论能够相互追溯。",
        14: "5. 试验设计与评价指标",
        15: "总体组合评价采用对象、网络、观测和随机种子全因子配对设计；主要比较RA-ECSC与云端时间戳预测补偿，关键次级比较为RA-ECSC与固定阈值边缘回退。另设置观测风险消融和滞回恢复消融，区分观测可信监督与切换机制的贡献。通信中断时长扫描用于识别边缘接管收益随中断持续时间的变化。",
        16: "主要终点为EC超出允许带的累计时间Tout；同时报告IAE、最大绝对误差、稳定时间、执行器总变差TV、切换次数和模式占用率。方法在相同随机日程下成对运行，以场景配对差为统计单位，采用种子块bootstrap给出区间，并报告中位数、P90及改善、近似和恶化比例。",
        17: "MQTT软件在环逐点核对风险分量、模式序列、转换、状态同步、降级锚点、预测量、EC和动作，并比较Tout、IAE和TV。图3概括对象、网络、观测、机制和软件实现组成的分层验证体系；实体台架只在传感器标定、执行联锁和人工接管完整后开展。",
        18: "图3  风险感知边云协同控制的分层验证体系",
    })

    set_paragraphs(method[1].cells[0].paragraphs, {
        1: "（1）将通信健康与本地EC观测可信度纳入统一监督器，形成面向水肥一体机的双风险三模式控制权分配方法，使边缘接管同时考虑网络可用性和反馈可靠性。",
        2: "（2）设计模型锚定降级控制与控制器状态同步机制，以最后可信观测修正模型预测，并结合滞回、驻留时间和恢复确认抑制异常期间的错误闭环与往返切换。",
        3: "（3）建立对象、网络、传感器和随机日程配对的冻结评价方法，通过主要对照、次级对照、机制消融和中断时长扫描识别控制收益及其适用条件。",
        4: "（4）把风险计算、模式转换、状态同步和任务审计落实到MQTT跨进程系统，实现参考算法与软件实现的逐点语义核对，为水肥装备的模块化部署提供可复现基础。",
    })

    set_paragraphs(method[2].cells[0].paragraphs, {
        2: "前期已完成RA-ECSC算法、冻结仿真协议和MQTT软件在环。1500个配对场景中，RA-ECSC相对云端时间戳预测补偿的累计超差时间配对中位数减少19.0 s，执行器总变差降低12.74%；相对固定阈值边缘回退的累计超差时间配对中位数减少17.3 s。通信中断扫描表明边缘接管收益随中断延长而增强；30条MQTT轨迹与MATLAB参考实现的模式、转换、同步和控制结果一致。",
        3: "研究已具备Python、MATLAB、Java、MQTT、数据库和前端开发基础，能够完成对象仿真、控制算法、故障代理、跨进程通信、任务审计和管理界面。现有结果可以支撑系统总体设计和软件实现，后续工作重点是统一接口、完善共同漂移处理、扩展执行器约束并形成可复现实验包。",
        5: "（1）名义模型与实际混肥动态存在偏差。通过对象参数族、运行中参数切换和独立辨识数据检验监督参数，模型残差阈值覆盖正常对象变化，避免把全部失配解释为传感故障。",
        6: "（2）双传感器共同漂移难以由通道差识别。增加独立参考EC、在线校准或与流量、配比相关的异构观测，在评价中单独报告共同漂移的中位结果与P90，不把总体优势外推到该工况。",
        7: "（3）模式切换可能引起执行器突变。采用PI积分状态反算、云边状态同步、陈旧命令清除、统一变化率限制和最小驻留时间，并通过切换次数与TV共同评价平滑性。",
        8: "（4）MQTT QoS 1允许重复且网络事件可能乱序。消息携带任务标识、逻辑序号和源时间戳，接收端执行幂等、陈旧拒绝和顺序检查；通过故障代理和重放试验验证消息语义。",
        9: "（5）平均改进可能掩盖尾部恶化。保持场景级配对，报告P90、改善比例和最差工况；控制器和监督阈值只在开发集选择，锁定评价后不追溯调参。",
        10: "（6）仿真与实体执行存在差异。台架验证先完成EC探头和注肥泵标定，再检查流量压力联锁、人工接管和异常暂停；实体结果与仿真结果分层报告。",
    })

    schedule = doc.tables[4]
    updates = {
        2: ("完成开题论证，冻结系统范围、控制对象、研究问题和分层验证协议。", "开题报告、需求规格和实验协议"),
        3: ("完善混肥对象辨识、公共PI、云端预测补偿和对象参数族。", "对象模型、底层控制器和验证数据"),
        4: ("实现通信风险、观测风险、三模式状态机和无扰切换，完成开发集参数冻结。", "RA-ECSC控制器与冻结配置"),
        5: ("开展总体组合评价、主要对照、机制消融和通信中断时长扫描。", "配对评价、区间估计和适用范围"),
        6: ("完成MQTT跨进程软件在环、消息异常与进程恢复验证。", "软件在环系统和一致性报告"),
        7: ("开发数据管理、任务审计、告警与可视化界面，完成系统集成。", "可运行原型、接口文档和管理界面"),
        8: ("开展共同漂移扩展、执行器友好约束和可选台架复核。", "扩展机制、台架记录和综合验证"),
        9: ("完成论文撰写、预审修改、答辩材料及代码数据归档。", "论文定稿、答辩材料和研究归档"),
    }
    for idx, (work, result) in updates.items():
        set_text(schedule.rows[idx].cells[1].paragraphs[0], work)
        set_text(schedule.rows[idx].cells[2].paragraphs[0], result)

    ref_cell = schedule.rows[15].cells[0]
    set_text(ref_cell.paragraphs[0], f"八、主要参考文献（共{len(REFERENCES)}项）")
    for i, paragraph in enumerate(ref_cell.paragraphs[1:]):
        set_text(paragraph, REFERENCES[i] if i < len(REFERENCES) else "")

    set_paragraphs(schedule.rows[16].cells[0].paragraphs, {
        1: "第1章 绪论：研究背景、国内外进展、研究问题、主要内容与技术路线。",
        2: "第2章 智能水肥一体化系统总体设计：需求、分层架构、数据对象、消息协议与任务流程。",
        3: "第3章 混肥EC对象建模与公共底层控制：对象辨识、参数变体、PI控制与执行器约束。",
        4: "第4章 风险感知边云协同控制方法：通信风险、观测风险、三模式状态机和无扰切换。",
        5: "第5章 通信与传感异常下的控制性能验证：配对场景、对照方法、消融、时长扫描与机制分析。",
        6: "第6章 MQTT软件在环与系统实现：跨进程架构、消息一致性、任务审计、管理界面与恢复机制。",
        7: "第7章 综合分析：适用工况、共同漂移边界、模型失配、执行器代价及实体扩展。",
        8: "第8章 结论与展望：系统实现、控制方法、验证结论和工程应用方向。",
    })

    normalize_document_fonts(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    with tempfile.TemporaryDirectory(prefix="proposal_t053_") as tmp:
        folder = Path(tmp)
        architecture = folder / "architecture.png"
        supervisor = folder / "supervisor.png"
        validation = folder / "validation.png"
        make_architecture(architecture); make_supervisor(supervisor); make_validation(validation)
        rebuild_with_media(OUTPUT, {
            "word/media/image2.png": architecture,
            "word/media/image3.png": supervisor,
            "word/media/image4.png": validation,
        })


if __name__ == "__main__":
    main()
