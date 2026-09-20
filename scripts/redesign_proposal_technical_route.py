from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw

from revise_proposal_from_current_evidence import (
    arrow,
    font,
    rebuild_with_media,
    set_text,
    wrapped,
)


SOURCE = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现_润色稿.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现_技术路线图重设计稿.docx"
)


def centered_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    *,
    fill: str,
    outline: str = "#506273",
    title: bool = False,
    radius: int = 22,
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=4)
    selected = font(30 if title else 27, bold=title)
    lines = wrapped(draw, text, selected, x2 - x1 - 44)
    line_height = 41 if title else 38
    y = y1 + (y2 - y1 - line_height * len(lines)) / 2
    for line in lines:
        bounds = draw.textbbox((0, 0), line, font=selected)
        width = bounds[2] - bounds[0]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=selected, fill="#1F2933")
        y += line_height


def phase_label(draw: ImageDraw.ImageDraw, y: int, number: str, label: str, color: str) -> None:
    draw.rounded_rectangle((45, y, 235, y + 105), radius=18, fill=color, outline="#506273", width=4)
    selected = font(28, bold=True)
    for offset, text in ((16, number), (57, label)):
        bounds = draw.textbbox((0, 0), text, font=selected)
        draw.text((140 - (bounds[2] - bounds[0]) / 2, y + offset), text, font=selected, fill="#1F2933")


def make_technical_route(
    path: Path,
    problem_text: str = "通信与传感异常并发条件下的水肥一体机EC连续控制",
) -> None:
    image = Image.new("RGB", (2200, 1680), "white")
    draw = ImageDraw.Draw(image)

    phase_label(draw, 55, "阶段1", "问题定义", "#DDEAF4")
    centered_box(
        draw,
        (300, 40, 2150, 175),
        problem_text,
        fill="#DDEAF4",
        title=True,
    )
    centered_box(draw, (300, 205, 900, 350), "控制对象与边界\n混肥过程  EC允许带  执行器约束", fill="#EEF4F8")
    centered_box(draw, (930, 205, 1530, 350), "异常与信息流\n时延丢包中断  偏置漂移尖峰", fill="#EEF4F8")
    centered_box(draw, (1560, 205, 2150, 350), "评价问题\n跟踪性能  动作代价  切换平滑性", fill="#EEF4F8")
    arrow(draw, (1225, 175), (1225, 205), width=5)

    phase_label(draw, 430, "阶段2", "系统建模", "#E6F1E8")
    centered_box(draw, (300, 405, 880, 590), "物联网管控架构\n感知执行—边缘—MQTT—云端—管理", fill="#E6F1E8", title=True)
    centered_box(draw, (910, 405, 1510, 590), "混肥对象与公共底层控制\n一阶惯性纯滞后模型  离散PI  动作约束", fill="#E6F1E8", title=True)
    centered_box(draw, (1540, 405, 2150, 590), "消息与任务模型\n时间戳  序号  幂等  版本与事件审计", fill="#E6F1E8", title=True)
    arrow(draw, (1225, 350), (1225, 405), width=5)

    phase_label(draw, 675, "阶段3", "控制方法", "#E8E4F3")
    centered_box(draw, (300, 650, 860, 825), "通信风险 R_net\n命令年龄  时延EWMA  窗口丢包率", fill="#E8E4F3", title=True)
    centered_box(draw, (1590, 650, 2150, 825), "观测风险 R_obs\n双EC差异  融合观测与模型残差", fill="#E8E4F3", title=True)
    centered_box(
        draw,
        (900, 635, 1550, 855),
        "RA-ECSC监督状态机\nM0云端预测  M1边缘控制  M2模型锚定降级\n滞回  驻留时间  恢复确认  状态同步",
        fill="#DCD5ED",
        title=True,
    )
    arrow(draw, (1225, 590), (1225, 635), width=5)
    arrow(draw, (860, 735), (900, 735), width=5)
    arrow(draw, (1590, 735), (1550, 735), width=5)

    phase_label(draw, 950, "阶段4", "分层验证", "#F8E8CF")
    centered_box(draw, (300, 925, 880, 1115), "配对仿真\n对象参数 × 网络异常 × 传感异常 × 随机种子", fill="#F8E8CF", title=True)
    centered_box(draw, (910, 925, 1510, 1115), "对照与机制分析\n云端基线  固定回退  风险消融  中断时长扫描", fill="#F8E8CF", title=True)
    centered_box(draw, (1540, 925, 2150, 1115), "MQTT软件在环\n消息异常  模式序列  状态同步  进程恢复", fill="#F8E8CF", title=True)
    arrow(draw, (1225, 855), (1225, 925), width=5)

    phase_label(draw, 1215, "阶段5", "评价收口", "#F4E1E4")
    centered_box(
        draw,
        (300, 1190, 2150, 1335),
        "性能与边界评价：Tout  IAE  最大误差  稳定时间  TV  切换次数  模式占用率  尾部工况",
        fill="#F4E1E4",
        title=True,
    )
    arrow(draw, (1225, 1115), (1225, 1190), width=5)
    centered_box(draw, (300, 1385, 880, 1595), "系统产出\n物联网边云协同管控原型\n接口与任务审计规范", fill="#EEF2F5", title=True)
    centered_box(draw, (910, 1385, 1510, 1595), "方法产出\n双风险三模式监督方法\n参数与适用条件", fill="#EEF2F5", title=True)
    centered_box(draw, (1540, 1385, 2150, 1595), "验证产出\n配对评价与软件在环报告\n失效边界与扩展接口", fill="#EEF2F5", title=True)
    arrow(draw, (1225, 1335), (590, 1385), width=5)
    arrow(draw, (1225, 1335), (1210, 1385), width=5)
    arrow(draw, (1225, 1335), (1845, 1385), width=5)

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(300, 300))


def main() -> None:
    document = Document(SOURCE)
    method_cell = document.tables[3].rows[0].cells[0]
    set_text(
        method_cell.paragraphs[17],
        "图3按照问题定义、系统建模、控制方法、分层验证和评价收口组织技术路线。仿真比较、机制分析与MQTT软件在环共享对象、异常日程和评价指标，使系统设计、控制逻辑与验证结论能够沿同一任务链追溯。",
    )
    set_text(method_cell.paragraphs[18], "图3  智能水肥一体化管控系统技术路线图")
    document.save(OUTPUT)

    with tempfile.TemporaryDirectory(prefix="proposal_t056_") as folder:
        technical_route = Path(folder) / "technical_route.png"
        make_technical_route(technical_route)
        rebuild_with_media(OUTPUT, {"word/media/image4.png": technical_route})


if __name__ == "__main__":
    main()
