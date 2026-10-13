from pathlib import Path
from copy import deepcopy
import re
import json
import shutil
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Volumes/out/Projects/开题报告/开题报告_基于物联网与协同控制的智能水肥一体化管控系统设计与实现_以红壤丘陵区为例_参考文献版.docx')
OUTPUT = Path('/Volumes/out/Projects/开题报告/开题报告_引用规范修订版.docx')

TEXTS = {
 (0,2): '精准农业将感知、数据处理和控制执行连接起来，为水肥装备的自动化管理提供系统基础[1-2]。随着物联网与农业自动化技术的发展，设备状态采集、远程管理和数据辅助决策成为智能水肥系统的重要组成部分[3-5]。本课题面向红壤丘陵区设施农业中生产单元分散的应用需求，关注传感节点、边缘控制器、云端服务和执行机构之间的状态与指令交换。水肥一体机通过注肥泵将母液注入灌溉水，并以电导率（electrical conductivity，EC）作为混合肥液浓度的主要反馈量；已有EC模型控制、系统辨识与模糊PID、自适应PI研究为混肥动态调节提供了方法依据[6-8]。水肥过程仿真及营养液自动配制研究还表明，控制设计需要结合输配过程、混合动态与EC、pH监测需求[9-10]。因此，本课题将设备互联、EC控制连续性和运行追溯作为系统设计的重点。',
 (0,3): '设备—边缘—云连续体为现场响应与远端计算的分工提供了架构基础[11-13]，边缘计算及多层工业控制研究进一步讨论了计算任务靠近现场部署的意义[14-15]。对于含时延的闭环，预测补偿与网络容错控制可用于处理通信不确定性，但控制效果仍受模型、网络状态和执行约束影响[16-17]。持续中断时，现场节点需要保持基本控制能力；若本地EC观测同时存在偏置、漂移或尖峰，直接转入边缘反馈仍可能产生错误控制动作。本课题据此联合评估通信状态与观测可信度，研究控制权、反馈来源和降级方式的协调机制。',
 (0,5): '本课题以红壤丘陵区设施农业为应用场景，拟设计基于物联网与协同控制的智能水肥一体化管控系统。系统由双EC传感与设备状态采集、注肥执行机构、边缘控制节点、云端控制服务、数据存储和管理界面组成，采用MQTT 3.1.1组织发布订阅通信[18]。控制层采用风险感知协同控制策略，在M0云端预测模式、M1边缘控制模式和M2模型锚定降级模式之间分配控制权。三种模式共享底层PI、饱和、抗积分饱和和动作变化率约束，监督器负责风险计算、模式转换、状态同步与恢复管理。',
 (1,2): '水肥装备研究已覆盖EC与pH监测、自动配制和远程管理。围绕混肥EC动态，基于EC模型的控制、系统辨识与模糊PID、自适应PI主要关注对象参数变化及设定值跟踪[6-8]；水肥过程仿真和营养液配制综述则从过程与系统层面梳理了控制需求[9-10]。自主水肥管理系统和基于物联网的微服务框架，将水源管理、数据服务与水肥执行连接起来[19-20]。物联网水肥管理与LoRaWAN灌溉系统研究进一步说明了农业现场感知、无线通信和决策执行的集成路径[21-22]。温室智能水肥研究涉及产量、品质及水分和养分利用效率评价[23]，远程EC调节研究则关注执行电机的抗扰控制[24]。这些工作分别提供了农艺评价、设备控制与系统集成依据；本课题进一步聚焦通信与观测异常并发时的控制权转移及反馈管理。',
 (1,4): '边缘计算为现场快速响应与云端模型服务之间的分层协作提供技术基础，多层工业控制研究讨论了云边计算资源与控制任务的配置[14-15]。预测补偿和网络容错控制可缓解时延、丢包等因素的影响[16-17]，但水肥一体机还需要将消息时间语义、状态同步和执行器约束纳入同一控制链。PID控制的过程模型、整定与实现方法为本地闭环提供基础[25]；模型预测控制则提供了预测与约束处理的方法框架[26]。在本课题中，云端承担时间戳预测补偿、模型管理和任务编排，边缘端承担本地PI闭环与监督执行，研究重点是依据通信和观测风险动态分配两端职责。',
 (1,5): '网络恢复后，云边积分状态差异和待执行的陈旧命令仍可能造成输出突变。抗积分饱和与无扰切换研究为控制器状态衔接提供了参考方法[27]。本课题拟将积分状态同步与消息清理结合，并设置滞回、最小驻留时间和连续恢复确认，评价这些机制对切换平滑性和控制恢复过程的影响。',
 (1,7): '模型诊断、农业物联网传感故障诊断和温室传感异常检测研究，为EC观测可信度评估提供了方法参考[28-30]。模型残差与功能冗余能够辅助判断异常，但残差也可能受到对象参数变化和模型失配影响[31-32]。多传感器融合及传感器健康状态转换研究，则为反馈通道选择和故障恢复提供了控制层依据[33-34]。营养液多传感监测研究进一步体现了多源信息在水肥过程监测中的应用[35]。本课题拟以双EC通道差与模型残差构建观测风险，作为控制权切换的依据；对于双通道共同漂移，单靠通道差难以识别，将单列工况并分析参考测量、校准或异构过程信息的辅助作用。',
 (1,9): '云—雾—边基础设施中的农业数字孪生部署研究，将模型与多层计算架构联系起来，为模型服务与现场系统的协作提供了参考[36]。结合上述研究，本课题关注水肥EC闭环中通信健康、观测可信度、控制权转移和软件实现之间的衔接，按照“双风险评估—三模式分配—无扰切换—MQTT实现—分层验证”组织研究，以统一的异常日程和执行约束评价方法及系统。',
}

def expand(s):
    nums=[]
    for term in s.split(','):
        if '-' in term:
            a,b=map(int,term.split('-')); nums.extend(range(a,b+1))
        else: nums.append(int(term))
    return nums

def main():
    doc=Document(SOURCE)
    for (ti,pi),text in TEXTS.items():
        p=doc.tables[ti].cell(0,0).paragraphs[pi]
        base=deepcopy(p.runs[0]._element.rPr) if p.runs and p.runs[0]._element.rPr is not None else None
        p.clear()
        for part in re.split(r'(\[[0-9,-]+\])',text):
            if not part: continue
            r=p.add_run(part)
            if base is not None: r._element.insert(0,deepcopy(base))
            r.font.superscript=bool(re.fullmatch(r'\[[0-9,-]+\]',part))
            if r.font.superscript: r.font.size=Pt(9)
    body=[]; seen_cells=set()
    for ti,t in enumerate(doc.tables):
        for ri,row in enumerate(t.rows):
            for c in row.cells:
                if c._tc in seen_cells: continue
                seen_cells.add(c._tc)
                if ti==4 and ri==15: continue
                body.extend(p.text for p in c.paragraphs)
    groups=[expand(m.group(1)) for p in body for m in re.finditer(r'\[([0-9,-]+)\]',p)]
    first=[]
    for g in groups:
        assert len(g)<=4,g
        for n in g:
            if n not in first: first.append(n)
    assert first==list(range(1,37)),first
    assert not any(re.search(r'\][\s]*\[',p) for p in body)
    refs=[p.text for p in doc.tables[4].rows[15].cells[0].paragraphs if re.match(r'^\[\d+\]',p.text)]
    assert len(refs)==36 and sum('[M]' in s for s in refs)==11
    assert 'RA-ECSC' not in '\n'.join(body)
    doc.save(OUTPUT)
    archive=ROOT/'output/proposal_revision/开题报告_引用规范修订版.docx'
    shutil.copy2(OUTPUT,archive)
    audit={'source':str(SOURCE),'output':str(OUTPUT),'reference_count':36,'book_count':11,'citation_groups':len(groups),'max_references_per_group':max(map(len,groups)),'first_appearance':first,'adjacent_groups':0,'edited_paragraphs':len(TEXTS)}
    a=ROOT/'output/researchwrite/t061_citation_format/audit.json'; a.parent.mkdir(parents=True,exist_ok=True)
    a.write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(audit,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
