from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt

from redesign_proposal_technical_route import make_technical_route
from revise_proposal_from_current_evidence import (
    normalize_document_fonts,
    rebuild_with_media,
    set_text,
)


SOURCE = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现_技术路线图插入优化稿.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与协同控制的智能水肥一体化管控系统设计与实现_以红壤丘陵区为例_最终题名稿.docx"
)
TITLE = "基于物联网与协同控制的智能水肥一体化管控系统设计与实现——以红壤丘陵区为例"


def main() -> None:
    document = Document(SOURCE)

    set_text(document.paragraphs[6], f"论文题目：{TITLE}")
    for run in document.paragraphs[6].runs:
        run.font.size = Pt(15.5)

    basis = document.tables[0].rows[0].cells[0].paragraphs
    set_text(
        basis[2],
        "红壤丘陵区设施农业地形起伏、生产单元分散，水肥设备需要在传感节点、边缘控制器、云端服务和执行机构之间持续交换状态与指令。水肥一体机通过注肥泵将母液按比例注入灌溉水，并以电导率（electrical conductivity，EC）作为混合肥液浓度的主要反馈量。混肥过程受到对象增益、惯性、纯滞后、母液浓度和流量波动影响；物联网链路中的时延、丢包、乱序、重复和通信中断又会进一步改变闭环行为。因此，区域应用需求需要同时落实到设备互联、控制连续性和运行追溯三个层面。",
    )
    set_text(
        basis[5],
        "本课题以红壤丘陵区设施农业为应用场景，拟设计基于物联网与协同控制的智能水肥一体化管控系统。系统由双EC传感与设备状态采集、注肥执行机构、边缘控制节点、MQTT通信、云端控制服务、数据存储和管理界面组成；控制层采用风险感知协同控制策略，在M0云端预测模式、M1边缘控制模式和M2模型锚定降级模式之间分配控制权。三种模式共享底层PI、饱和、抗积分饱和和动作变化率约束，监督器负责风险计算、模式转换、状态同步与恢复管理。",
    )
    set_text(
        basis[8],
        "本研究将红壤丘陵区的分散生产单元和网络波动作为系统设计场景，关注水肥装备中的控制权管理问题。方法层面把命令时效性与本地观测可信度纳入同一决策过程，使云端计算、边缘响应和异常降级具有明确的切换依据；系统层面统一记录风险、模式、同步、命令和反馈，为控制过程复核、异常定位和后续区域化设备接入建立可追溯的数据基础。",
    )

    objectives = document.tables[2].rows[0].cells[0].paragraphs
    set_text(
        objectives[2],
        "总体目标是面向红壤丘陵区设施农业应用，设计并实现基于物联网与协同控制的智能水肥一体化管控系统，形成从双EC与设备状态采集、混肥对象建模、云端预测、边缘控制、模型锚定降级、风险监督和无扰切换，到MQTT通信与任务审计的完整技术链。研究将通过预先规定的仿真与软件在环协议，评价通信异常与传感异常并发条件下的EC调节性能、执行器动作代价和模式切换平滑性，并界定系统适用范围。",
    )
    set_text(
        objectives[6],
        "（1）系统总体设计。结合红壤丘陵区生产单元分散、通信链路波动和设备运行条件差异，划分感知与执行、边缘控制、通信、云端服务、数据管理和用户交互模块；规定EC、设备状态、控制命令、模式事件和执行反馈的数据结构，明确模块接口、消息方向和任务流程。",
    )

    method = document.tables[3].rows[0].cells[0].paragraphs
    set_text(
        method[2],
        "研究按照“区域应用约束与对象定义—混肥模型—公共底层PI—通信与观测风险—三模式监督—无扰切换—配对仿真—MQTT软件在环—系统集成”的顺序推进。红壤丘陵区在本研究中作为系统需求和工况组织的应用场景，各比较方法共享对象日程、底层控制器和执行器约束，监督策略仅改变控制位置、反馈来源与降级逻辑。",
    )

    normalize_document_fonts(document)
    document.save(OUTPUT)

    with tempfile.TemporaryDirectory(prefix="proposal_t058_") as folder:
        route = Path(folder) / "technical_route.png"
        make_technical_route(
            route,
            problem_text="红壤丘陵区应用场景下通信与传感异常并发的水肥一体机EC连续控制",
        )
        rebuild_with_media(OUTPUT, {"word/media/image4.png": route})


if __name__ == "__main__":
    main()
