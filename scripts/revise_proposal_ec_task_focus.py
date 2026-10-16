"""Revise the proposal around EC control, task management and system implementation."""
from pathlib import Path
from copy import deepcopy
import hashlib,json,re,shutil
from zipfile import ZipFile
from docx import Document
from docx.shared import Mm,Pt
from docx.oxml.ns import qn
from lxml import etree
import illustrate_proposal_portable as figures
import matplotlib as mpl
from revise_proposal_portable_prototype import set_text

mpl.rcParams.update({'font.family':'sans-serif','font.sans-serif':[figures.font_manager.FontProperties(fname=str(figures.FONT)).get_name(),'Arial','DejaVu Sans'],'font.size':8,'svg.fonttype':'none','pdf.fonttype':42})

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('/Volumes/out/Projects/开题报告/开题报告_便携样机图文完善版.docx')
OUTPUT=Path('/Volumes/out/Projects/开题报告/开题报告_EC调节与接入任务主线修订版.docx')
ASSET=ROOT/'output/researchwrite/t064_ec_task_focus'
REPLACEMENTS={
 (0,0,2):'精准农业将感知、数据处理和控制执行连接起来，为水肥装备的自动化管理提供系统基础[1-2]。物联网与农业自动化技术进一步将设备状态采集、远程管理和数据辅助决策纳入水肥作业流程[3-5]。本课题面向红壤丘陵区分散生产单元，以智能水肥一体化管控系统为研究主体，重点研究混肥EC调节与接入任务管理。电导率（electrical conductivity，EC）是混合肥液调节的反馈量；EC模型控制、系统辨识与模糊PID、自适应PI为动态调节提供了方法依据[6-8]。输配过程仿真和营养液配制研究也说明，控制设计应结合管路输配、混合动态及EC、pH监测需求[9-10]。系统设计需要把测量、注肥动作和任务配置联系起来，使EC控制与设备作业过程保持一致。',
 (0,0,3):'设备—边缘—云连续体为现场响应与远端计算的分工提供了架构基础[11-13]，边缘计算及多层工业控制研究讨论了控制任务靠近现场部署的方法[14-15]。预测补偿和网络容错控制为通信不确定性处理提供了依据[16-17]。水肥作业中的EC调节既依赖消息及时到达，也依赖现场观测可信；任务重新启动时，还需要确认配方、接入配置与控制状态。由此，本课题将异常期间的控制权分配与新任务的状态初始化纳入同一系统方案。',
 (0,0,5):'本课题拟设计基于物联网与协同控制的智能水肥一体化管控系统，围绕EC调节、接入任务、远程监视和作业追溯开展研究。系统由状态采集、注肥执行接口、边缘控制、云端服务、数据库和管理界面构成，采用MQTT 3.1.1组织发布订阅通信[18]。控制层联合通信风险与观测风险，在M0云端预测、M1边缘控制和M2模型锚定降级之间分配控制权；任务层关联设备、接入点、配方与模型版本，管理启动检查、命令执行及反馈。移动式水肥一体机作为实体载体，其多接入点轮换作业用于检验任务交接与控制恢复机制。',
 (0,0,6):'研究按“系统需求与总体设计—混肥建模与EC控制—接入任务管理—平台实现与验证”展开。算法对照评价正常和异常条件下的EC跟踪与动作平滑性，系统试验核对任务、命令和反馈的一致性；移动接入场景进一步检查重新接入后的状态初始化与EC恢复。',
 (0,0,8):'本课题将混肥EC控制与物联网作业管理连接起来，使控制模式、任务配置和执行反馈具有对应关系。研究拟以通信与观测双风险协调云边职责，以任务标识和状态初始化衔接设备启动与连续作业，并通过平台保存可追溯记录。面向红壤丘陵区分散地块，移动装备的轮换接入为上述机制提供应用场景，研究价值体现在EC调节、任务交接和系统运行的协同设计。',
 (1,0,9):'农业数字孪生部署研究将模型服务与云—雾—边基础设施联系起来，为现场系统与多层计算的协作提供参考[36]。移动式果园灌溉施肥机研究已集成混肥、检测、自动控制及牵引装置，并开展灌溉施肥性能试验[37]。本课题围绕EC连续调节和接入任务管理，研究通信与观测风险下的控制权分配，以及任务、命令与控制状态的衔接。移动式装备用于拓展多接入点应用：转场后重新核对配置、建立新任务和可信控制状态，再恢复EC调节。图1为文献中的牵引式移动装备示例。',
 (2,0,2):'面向红壤丘陵区水肥作业需求，设计并实现基于物联网与协同控制的智能水肥一体化管控系统，形成混肥EC调节、接入任务管理、远程监视和作业追溯的功能链。重点研究通信与观测异常下的EC协同控制，以及任务启动和命令执行的一致性。通过配对仿真、MQTT软件在环和样机对照评价控制性能与系统功能，并以移动装备的不同接入点切换检验应用适配性。',
 (2,0,3):'系统由感知执行、边缘控制、MQTT通信、云端服务、数据库和管理界面构成。边缘节点采集EC与设备状态，执行风险监督和本地控制；云端提供预测、配方模型管理及任务服务。接入任务关联设备、接入点、目标EC和参数版本，观测、命令及反馈均归入对应任务。移动作业沿用这一架构，通过新的接入配置和任务初始化完成地块间交接。总体架构如图2所示。',
 (2,0,5):'图2  智能水肥一体化管控系统总体架构（自绘）',
 (2,0,7):'（1）系统需求分析与总体设计。围绕EC设定与调节、数据采集、设备接入、任务启动、异常告警和历史追溯，确定功能模块与接口。划分云端、边缘和设备端职责，设计设备—接入点—任务的数据关系、命令反馈流程及状态管理规则，形成系统架构和功能验收清单。',
 (2,0,8):'（2）混肥对象建模与EC协同控制。完成EC、流量和注肥输出标定，采集入口水EC、出口EC、温度、流量压力及注肥动作，以不同完整试验分别进行模型辨识和验证。建立一阶惯性纯滞后模型与公共离散PI，联合通信风险和观测风险设计M0、M1、M2监督，以及状态同步、滞回和恢复确认机制。',
 (2,0,9):'（3）接入任务管理与控制状态衔接。设计接入点配置、配方与模型版本核对、任务建立、启动检查和执行反馈。新任务清理旧命令、PI状态和降级锚点，建立本次作业的可信状态；同一任务切换控制模式时实施积分反算和云边同步。移动装备转场后沿用新任务流程，检验跨接入点状态隔离与启动恢复。',
 (2,0,10):'（4）物联网管控平台实现。采用Python开发模型、控制与评价模块，围绕样机控制器适配采集和执行接口，以MQTT 3.1.1与QoS 1连接云边进程。实现实时监视、接入配置、任务启停、异常告警与历史查询；数据库关联观测、风险、模式、命令、反馈及参数版本，支持作业过程追溯。',
 (2,0,11):'（5）控制性能与系统功能验证。以配对仿真检验EC控制、基线差异及关键机制，以软件在环检验消息和任务一致性，以样机试验评价EC启动、跟踪和短时通信中断恢复。移动接入作为专项场景，比较不同接入点切换后的配置检查、旧命令拒绝和EC恢复过程。',
 (2,0,13):'（1）如何依据混肥动态与观测可信度，在通信和传感异常下协调云边控制，并保持EC调节的连续性。',
 (2,0,14):'（2）如何区分同一任务内的控制模式切换与新任务间的状态重建，使接入配置、控制状态和执行命令保持一致。',
 (2,0,15):'（3）如何在共同对象模型、底层PI和执行约束下，评价EC跟踪、异常恢复与执行器动作之间的关系。',
 (2,0,16):'（4）如何通过设备—接入点—任务关联实现作业追溯，并适配移动装备的重新接入。',
 (3,0,2):'技术路线围绕系统实现、EC调节和接入任务管理展开。首先确定系统需求、数据关系及感知执行接口，开展混肥动态测量与独立模型验证；随后设计双风险三模式EC控制、无扰切换和任务启动管理；再实现MQTT通信、数据库与管理界面，组织仿真、软件在环和样机验证。移动式装备的多接入点切换纳入应用验证，重点检查新任务初始化和EC恢复。',
 (3,0,4):'图3  智能水肥一体化管控系统研究技术路线图',
 (3,0,5):'2. 混肥对象建模与EC调节',
 (3,0,6):'对象建模先完成EC探头、流量及注肥输出标定，再开展阶跃、启动和设定值变化试验。同步记录入口水EC、母液种类与浓度、温度、出口EC、流量压力和控制动作，以独立的完整运行记录分别进行辨识和验证。模型以增益、时间常数和纯滞后描述混肥响应，验证残差用于观测风险校准。采样周期依据响应速度及接口能力确定，并与控制参数在正式评价前固定。',
 (3,0,7):'三种模式共用并联离散PI、饱和、抗积分饱和和动作变化率限制。M0利用时间戳与名义模型补偿观测时延，M1以可信边缘观测闭环，M2以同一任务内最后可信观测校正模型递推。各对照方法使用相同对象工况、模型和底层参数，比较重点为风险监督及模式衔接对EC调节的作用。',
 (3,0,10):'通信风险超过阈值且观测可信时，由M0转入M1；观测风险升高时进入M2。滞回、最小驻留时间和连续恢复确认管理模式恢复，PI积分反算与云边同步衔接切换动作。任务启动则执行独立的初始化流程：停止注肥，清理旧命令、积分状态与模型锚点，核对接入配置与配方，通水并确认可信观测后从本地控制启动，再按同步条件进入M0。降级锚点由本次作业建立。',
 (3,0,11):'接入任务管理覆盖配置核对、任务建立、启动、运行和结束。设备、接入点、配方、模型及参数版本在启动前绑定；命令与反馈携带任务标识，接收端拒绝跨任务、重复或陈旧命令，并记录执行状态。同一任务内保存模式切换与状态同步事件，新任务重新建立控制状态。移动装备更换接入点时，按这一流程组织任务交接。任务管理与EC控制的衔接如图4所示。',
 (3,0,13):'图4  接入任务管理与EC协同控制流程（自绘）',
 (3,0,16):'平台界面按实时监视、接入配置、任务管理、告警和历史查询组织。运行页展示EC设定值与反馈、流量压力、风险和控制模式；任务页管理配方、接入点、启停状态与执行反馈；历史页按任务关联数据、命令和事件。数据库与消息记录使用一致的任务及版本标识，使界面显示、控制执行和作业追溯相互对应。',
 (3,0,18):'控制算法采用对象工况、网络状态、观测状态和随机种子配对设计，比较云端时间戳预测、固定阈值边缘回退与双风险协同控制。正常场景评价EC跟踪与动作代价，异常场景评价通信中断、观测偏差及恢复过程。分别消融观测风险与切换机制，检验反馈选择及状态同步的作用。模型、底层PI和风险参数在开发阶段确定，独立评价后不再调整。',
 (3,0,19):'样机试验以EC调节为核心，先开展正常启动、设定值跟踪与稳态作业，再实施可控短时通信中断和重连。低、中、高三个供水压力水平用于表征对象工况，具体值在设备额定范围内经预试验确定；各方法采用共同配方、初始状态和执行约束。每种方法在同一工况下初步安排不少于3次独立启动—作业重复，按批次轮换方法顺序，统一通水、清洗、母液配制及初始EC确认流程。最终重复数依据预试验变异和主要终点精度确定。',
 (3,0,20):'EC控制以超出允许带的累计时间Tout为主要指标，辅以IAE、最大误差、稳定时间、执行器总变差TV、切换次数和模式占用率。任务管理评价命令反馈对应、旧任务命令拒绝、状态重建及事件追溯；移动接入专项增加接入后EC恢复时间和启动成功率。EC允许带、持续达标判据、计时起点与功能通过条件在正式试验前固定，失败记录纳入汇总。以完整运行作为重复单位，配对报告差值、区间及尾部工况，分层解释仿真、软件在环和样机结果。',
 (3,1,1):'（1）通信与观测双风险驱动的EC协同控制。拟依据命令时效性和现场EC观测可信度，在云端预测、边缘控制与模型锚定降级之间分配控制权，并通过积分反算与状态同步衔接模式切换。以共同底层PI下的配对对照和机制消融，检验异常期间EC调节与动作平滑性。',
 (3,1,2):'（2）接入任务与控制状态的协同管理。拟将设备、接入点、配方及模型版本关联到具体任务，区分跨任务状态初始化与同一任务内的无扰切换，使启动检查、命令执行和控制恢复按统一任务关系运行。以跨任务命令拒绝、反馈一致性和启动对照验证这一衔接机制。',
 (3,1,3):'（3）面向红壤丘陵区分散地块的移动接入应用特色。以移动式水肥一体机为载体，将不同接入点的配置切换、旧状态清理及EC恢复纳入平台流程。通过重新接入专项试验，检验系统在地块间轮换作业时的任务交接与控制恢复能力。',
 (3,1,4):'课题以EC协同控制和接入任务管理为研究重点，以物联网管控系统实现连接算法与作业功能，移动接入体现面向分散地块的应用特色。',
 (3,2,2):'本课题以水肥一体机及其感知执行接口作为实体试验载体，开展EC与注肥标定、混肥动态辨识、闭环对照和任务管理验证。本地计算环境承担对象仿真、异常注入、统计评价及MQTT软件在环；移动式样机用于进一步检验不同接入点下的配置管理、启动与控制恢复。',
 (3,2,5):'（1）混肥动态与模型配置随供水和管路工况变化。先开展额定范围内的响应试验，分离辨识与验证数据；用独立验证残差校准观测风险，在模型失配场景中检查误切换，并为接入配置绑定适用模型。',
}

TASK_TEST='系统功能验证覆盖设备接入、任务启停、告警、历史查询及命令反馈关联，通过重复、乱序、延迟和跨任务消息注入核对处理结果。移动接入专项初步选择两个接入点配置，统一配方和目标EC，比较规范初始化与原有启动流程；记录从接入检查通过到EC持续进入允许带的时间，核对旧命令隔离和任务状态重建。接入点条件按样机接口及预试验确定。'

def make_figures():
 ASSET.mkdir(parents=True,exist_ok=True)
 figures.ASSET=ASSET
 f,a=figures.canvas(153)
 # Original design schematic: three study threads, mobility is an application scenario.
 stages=[(123,20,'需求与架构'),(94,22,'对象与任务'),(64,23,'协同方法'),(32,24,'分层验证'),(3,21,'评价与实现')]
 for y,h,title in stages:figures.box(a,2,y,24,h,title,figures.GRAY,size=8)
 for x,s in [(30,'系统功能\n监视 / 告警 / 追溯'),(74,'EC调节\n目标 / 观测 / 执行'),(118,'移动接入场景\n配置切换 / 启动恢复')]:figures.box(a,x,123,40,20,s,figures.BLUE,size=7)
 figures.box(a,30,94,62,22,'混肥动态与公共PI\n标定 / 辨识 / 独立验证',figures.TEAL,size=8)
 figures.box(a,98,94,60,22,'设备—接入点—任务\n配置 / 版本 / 命令反馈',figures.TEAL,size=8)
 figures.box(a,30,64,62,23,'双风险三模式EC控制\nM0云端 / M1边缘 / M2降级\n积分反算 / 状态同步',figures.BLUE,size=7.3)
 figures.box(a,98,64,60,23,'接入任务与状态管理\n检查 / 初始化 / 运行 / 结束\n旧命令隔离 / 执行追溯',figures.BLUE,size=7.3)
 for x,s in [(30,'配对仿真\n基线 / 异常 / 机制消融'),(74,'平台与软件在环\n功能 / 消息 / 任务一致性'),(118,'样机EC对照\n跟踪 / 中断 / 移动接入')]:figures.box(a,x,32,40,24,s,figures.TEAL,size=7)
 figures.box(a,30,3,128,21,'EC性能：Tout / IAE / 稳定时间 / TV\n系统功能：任务一致性 / 告警 / 追溯   移动专项：接入后恢复\n形成智能水肥一体化管控系统、控制方法与分层验证报告',figures.GRAY,size=7.2)
 for upper,lower in [(123,116),(94,87),(64,56),(32,24)]:figures.arrow(a,(94,upper),(94,lower))
 f.savefig(ASSET/'technical_route_ec_task.png',dpi=600,facecolor='white')
 f.savefig(ASSET/'technical_route_ec_task.svg',facecolor='white')
 f.savefig(ASSET/'technical_route_ec_task.pdf',facecolor='white')
 figures.plt.close(f)
 # Reuse the existing Python design and update only the scope labels of Fig. 4.
 src=Path(figures.__file__).read_text().split('def after(')[0]
 src=src.replace("ASSET=ROOT/'output/researchwrite/t063_illustrated_proposal'",f"ASSET=Path({str(ASSET)!r})")
 src=src.replace('a  转场后的新任务初始化','a  接入任务建立与启动')
 ns={'__name__':'t064_flow_design','__file__':str(Path(figures.__file__))}
 exec(compile(src,'t064_flow_design','exec'),ns);ns['figures']()
 # Unchanged photo and architecture remain archived in T063; keep this revision focused.
 for name in ['mobile_machine_reference.png','system_architecture.png','system_architecture.svg','system_architecture.pdf']:
  (ASSET/name).unlink(missing_ok=True)

def replace_media(doc,paragraph,path):
 rid=paragraph._p.xpath('.//a:blip/@r:embed')[0]
 doc.part.related_parts[rid]._blob=path.read_bytes()

def main():
 make_figures();d=Document(SOURCE)
 refs_before=d.tables[4].rows[15].cells[0].text
 # Use the user's broader term; no chassis or lightweight-design claim is added.
 seen=set()
 for t in d.tables:
  for row in t.rows:
   for c in row.cells:
    if c._tc in seen:continue
    seen.add(c._tc)
    for p in c.paragraphs:
     for r in p.runs:
      if any(s in r.text for s in ['便携式','便携样机','便携水肥']):r.text=r.text.replace('便携式','移动式').replace('便携样机','移动样机').replace('便携水肥','移动水肥')
 for (ti,ri,pi),text in REPLACEMENTS.items():set_text(d.tables[ti].rows[ri].cells[0].paragraphs[pi],text)
 d.tables[0].cell(0,0).paragraphs[4].paragraph_format.keep_with_next=True
 d.tables[0].cell(0,0).paragraphs[5].paragraph_format.widow_control=True
 d.tables[1].cell(0,0).paragraphs[0].paragraph_format.page_break_before=True
 d.tables[2].cell(0,0).paragraphs[6].paragraph_format.keep_with_next=True
 d.tables[2].cell(0,0).paragraphs[6].paragraph_format.page_break_before=True
 d.tables[2].cell(0,0).paragraphs[7].paragraph_format.widow_control=True
 method=d.tables[3].cell(0,0)
 figures.after(method.paragraphs[19],TASK_TEST)
 added=method.paragraphs[20]
 added._p.remove(added._p.pPr)
 if method.paragraphs[19]._p.pPr is not None:added._p.insert(0,deepcopy(method.paragraphs[19]._p.pPr))
 added.alignment=0;added.paragraph_format.first_line_indent=Mm(7.4)
 for r in added.runs:
  if r._r.rPr is not None:r._r.remove(r._r.rPr)
  if method.paragraphs[19].runs[0]._r.rPr is not None:r._r.insert(0,deepcopy(method.paragraphs[19].runs[0]._r.rPr))
 route=method.paragraphs[3]
 route.clear();route.paragraph_format.line_spacing=1;route.paragraph_format.first_line_indent=Mm(0)
 route.paragraph_format.left_indent=Mm(0);route.paragraph_format.right_indent=Mm(0)
 route.paragraph_format.keep_with_next=True;route.add_run().add_picture(str(ASSET/'technical_route_ec_task.png'),width=Mm(132))
 # Paragraph indices after the new experiment paragraph are unchanged for the flow figure.
 replace_media(d,method.paragraphs[12],ASSET/'task_control_flow.png')
 schedule={3:('完成设备采集执行接口和EC注肥标定，开展混肥辨识与独立验证，建立任务及接入数据模型。','接口规格、混肥模型与任务模型'),4:('实现双风险三模式EC控制与无扰切换，设计任务启动检查、状态初始化及命令隔离。','EC协同控制与任务管理方法'),5:('开展正常与异常条件配对仿真、控制基线及机制消融，固定独立评价协议。','EC控制评价与机制分析'),6:('实现MQTT消息链和任务服务，核对消息处理、命令反馈与控制状态一致性，适配样机接口。','软件在环报告与设备适配模块'),7:('完成实时监视、任务启停、告警及历史查询界面，开展功能验收与样机预试验。','管控系统与对照试验协议'),8:('开展样机EC跟踪和通信中断恢复对照，完成接入任务核对与移动接入专项验证。','EC与系统功能数据及专项评价')}
 for ri,(content,result) in schedule.items():
  set_text(d.tables[4].rows[ri].cells[1].paragraphs[0],content);set_text(d.tables[4].rows[ri].cells[2].paragraphs[0],result)
 outline=['九、论文详细提纲','第1章 绪论：水肥管控需求、EC调节与协同控制进展、红壤丘陵区应用场景。','第2章 智能水肥一体化管控系统总体设计：功能需求、设备接口、云边架构和任务数据关系。','第3章 混肥对象建模与EC控制基础：标定、辨识、独立验证及公共PI。','第4章 风险感知EC协同控制方法：双风险监督、三模式分配、状态同步与无扰切换。','第5章 接入任务管理与平台实现：配置核对、任务启停、命令反馈、告警与追溯。','第6章 控制性能与系统功能验证：配对仿真、软件在环、样机EC对照及移动接入专项。','第7章 综合分析：EC调节与动作代价、任务一致性、接入恢复及模型适用条件。','第8章 结论与展望：系统实现、EC协同控制、任务管理与应用扩展。']
 for p,text in zip(d.tables[4].rows[16].cells[0].paragraphs,outline):set_text(p,text)
 assert d.tables[4].rows[15].cells[0].text==refs_before
 OUTPUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUTPUT)
 shutil.copy2(OUTPUT,ROOT/'output/proposal_revision'/OUTPUT.name)
 body=[];seen=set()
 for ti,t in enumerate(d.tables):
  for ri,row in enumerate(t.rows):
   for c in row.cells:
    if c._tc in seen or (ti==4 and ri==15):continue
    seen.add(c._tc);body.extend(p.text.replace('‑','-') for p in c.paragraphs)
 order=[];maxgroup=0
 for match in re.finditer(r'\[([0-9,-]+)\]','\n'.join(body)):
  vals=[]
  for v in match[1].split(','):
   if '-' in v:
    a,b=map(int,v.split('-'));vals.extend(range(a,b+1))
   else:vals.append(int(v))
  maxgroup=max(maxgroup,len(vals))
  for v in vals:
   if v not in order:order.append(v)
 assert order==list(range(1,38));assert maxgroup<=4
 text='\n'.join(body)
 for forbidden in ['RA-ECSC','项目未开始','我要参加','已取得成果','便携式','轻量化设计','行走机构设计']:assert forbidden not in text,forbidden
 with ZipFile(OUTPUT) as z:assert z.testzip() is None
 audit={'output':str(OUTPUT),'modified_paragraphs':len(REPLACEMENTS),'added_experiment_paragraphs':1,'references_unchanged':True,'references':37,'books':11,'first_appearance':order,'max_citations_per_group':maxgroup,'figures':4,'redrawn_figures':[3,4],'unchanged_photo_source':'文献37原图5，示例属性不变','mainline':'物联网管控系统实现、EC协同控制与接入任务管理；移动接入为应用特色','new_experimental_results':False}
 for p in ASSET.glob('*.svg'):p.write_text('\n'.join(line.rstrip() for line in p.read_text().splitlines())+'\n')
 (ASSET/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n');print(json.dumps(audit,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
