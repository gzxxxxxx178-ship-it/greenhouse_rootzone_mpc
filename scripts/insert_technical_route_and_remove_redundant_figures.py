from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from redesign_proposal_technical_route import make_technical_route
from revise_proposal_from_current_evidence import normalize_document_fonts, set_text


SOURCE = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现_技术路线图重设计稿.docx"
)
OUTPUT = Path(
    "/Volumes/out/Projects/python/greenhouse_rootzone_mpc/output/proposal_revision/"
    "开题报告_基于物联网与边云协同控制的智能水肥一体化管控系统设计与实现_技术路线图插入优化稿.docx"
)


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    return Paragraph(element, paragraph._parent)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def remove_unused_image_relationships(document: Document) -> None:
    for relationship_id in list(document.part.rels):
        relationship = document.part.rels[relationship_id]
        if "image" not in relationship.reltype:
            continue
        if relationship.target_ref in {
            "media/image2.png",
            "media/image3.png",
            "media/image4.png",
        }:
            del document.part.rels[relationship_id]


def main() -> None:
    document = Document(SOURCE)

    overview = document.tables[2].rows[0].cells[0]
    set_text(
        overview.paragraphs[3],
        "系统总体架构由边缘感知与执行、MQTT通信、云端服务、风险监督、数据存储和管理界面构成。边缘侧负责双EC、流量、压力、液位和执行器状态的时间戳标记、质量检查与预处理；云端负责预测补偿、模型版本管理和任务编排；监督器依据通信风险与观测风险选择M0、M1或M2。数据库记录观测、风险、模式、命令、反馈和版本，管理界面呈现实时状态、告警与历史任务。",
    )
    remove_paragraph(overview.paragraphs[4])

    method_cell = document.tables[3].rows[0].cells[0]
    set_text(
        method_cell.paragraphs[9],
        "三模式监督器不另行构造底层控制律，而是确定当前控制位置、反馈来源和降级机制。每次模式转换记录触发风险、来源模式、目标模式、同步状态和最终动作，为切换过程分析和软件语义核对提供事件依据。",
    )
    remove_paragraph(method_cell.paragraphs[10])

    # Removing the earlier caption changes subsequent indices by one.
    method_cell = document.tables[3].rows[0].cells[0]
    set_text(
        method_cell.paragraphs[16],
        "MQTT软件在环将逐点核对风险分量、模式序列、转换事件、状态同步、降级锚点、预测量、EC和控制动作，并比较Tout、IAE和TV等汇总指标。实体台架验证仅在EC探头标定、注肥执行联锁和人工接管流程明确后作为扩展环节开展。",
    )
    remove_paragraph(method_cell.paragraphs[17])

    remove_unused_image_relationships(document)

    with tempfile.TemporaryDirectory(prefix="proposal_t057_") as folder:
        route_path = Path(folder) / "technical_route.png"
        make_technical_route(route_path)

        method_cell = document.tables[3].rows[0].cells[0]
        anchor = method_cell.paragraphs[2]
        image_paragraph = paragraph_after(anchor)
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(4)
        image_paragraph.paragraph_format.space_after = Pt(2)
        image_paragraph.add_run().add_picture(str(route_path), width=Inches(5.25))

        caption = paragraph_after(image_paragraph)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(6)
        caption_run = caption.add_run("图1  智能水肥一体化管控系统技术路线图")
        caption_run.font.size = Pt(10.5)

        normalize_document_fonts(document)
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        document.save(OUTPUT)


if __name__ == "__main__":
    main()
